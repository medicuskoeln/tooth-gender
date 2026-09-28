import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.patches as mpatches
from matplotlib.patches import FancyBboxPatch
from matplotlib.colors import Normalize
from matplotlib import cm
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
import os, tempfile

# =====================================================================
# DATA LOADING & MATCHING (same as Tab 3, WA/FA/OA separate)
# =====================================================================
df = pd.read_excel('roh1.xlsx')
df = df[df['Geschlecht_w1_m2'].isin([1, 2])]

# Filter SLT only (tubus == 1)
df = df[df['tubus'] == 1]
print(f"SLT only: {len(df)} rows")

matching_cols = ['Ausbild_WA1_FA2_OA3', 'Laryng_diff_1Mac_2Glide_3King_4CMAC', 'tubus', 'atemweg']

matched_pairs = []
for _, group in df.groupby(matching_cols):
    w = group[group['Geschlecht_w1_m2'] == 1]
    m = group[group['Geschlecht_w1_m2'] == 2]
    n = min(len(w), len(m))
    if n > 0:
        matched_pairs.append((
            w.sample(n=n, random_state=42).reset_index(drop=True),
            m.sample(n=n, random_state=42).reset_index(drop=True)
        ))

df_w = pd.concat([w for w, _ in matched_pairs], ignore_index=True)
df_m = pd.concat([m for _, m in matched_pairs], ignore_index=True)
print(f"Matched pairs (SLT): {len(df_w)}")

laryng = {1: 'Macintosh', 2: 'GlideScope', 3: 'King Vision', 4: 'C-MAC'}
teeth_cols = [
    ('res_12_max', '12'),
    ('res_11_max', '11'),
    ('res_21_max', '21'),
    ('res_22_max', '22'),
]

# =====================================================================
# COMPUTE MEDIANS per laryngoscope × airway × sex × tooth
# =====================================================================
data = {}  # (aw, lv, sex) -> {tooth_label: median}

for aw, aw_label in [(1, 'standard'), (2, 'difficult')]:
    for lv in laryng:
        for sex, sex_label, df_sex in [(1, 'female', df_w), (2, 'male', df_m)]:
            sub = df_sex[(df_sex['atemweg'] == aw) &
                         (df_sex['Laryng_diff_1Mac_2Glide_3King_4CMAC'] == lv)]
            n = len(sub)
            vals = {}
            for col, label in teeth_cols:
                v = pd.to_numeric(sub[col], errors='coerce').dropna()
                vals[label] = v.median() if len(v) > 0 else 0
            data[(aw, lv, sex_label)] = {'vals': vals, 'n': n}
            print(f"  {aw_label} | {laryng[lv]:12s} | {sex_label:6s} | n={n:3d} | "
                  f"12={vals['12']:.1f}  11={vals['11']:.1f}  21={vals['21']:.1f}  22={vals['22']:.1f}")

# Global min/max for consistent colorbar
all_medians = []
for k, v in data.items():
    all_medians.extend(v['vals'].values())
vmin = 0
vmax = max(all_medians) * 1.05
print(f"\nColor scale: {vmin:.1f} – {vmax:.1f} N")

# =====================================================================
# DRAW TOOTH FUNCTION
# =====================================================================
cmap = cm.YlOrRd  # yellow → orange → red
norm = Normalize(vmin=vmin, vmax=vmax)

def draw_upper_jaw(ax, tooth_vals, title='', n=0):
    """Draw a simplified frontal view of upper 4 front teeth, colored by force.
    tooth_vals: dict with keys '12','11','21','22' -> median force value
    """
    ax.set_xlim(-2.5, 2.5)
    ax.set_ylim(-0.5, 3.5)
    ax.set_aspect('equal')
    ax.axis('off')

    # Gum line (arc)
    theta = np.linspace(0.15 * np.pi, 0.85 * np.pi, 100)
    gum_x = 2.8 * np.cos(theta)
    gum_y = 2.8 * np.sin(theta) + 0.3
    ax.fill_between(gum_x, gum_y, 3.5, color='#F5C6C6', alpha=0.4, zorder=0)
    ax.plot(gum_x, gum_y, color='#D4738C', lw=1.2, zorder=1)

    # Tooth positions and widths (frontal view, slight arch)
    # Order: 12(right lateral), 11(right central), 21(left central), 22(left lateral)
    tooth_info = {
        '12': {'x': -1.45, 'w': 0.7, 'h': 1.5, 'gy': 1.9},   # right lateral incisor
        '11': {'x': -0.55, 'w': 0.85, 'h': 1.8, 'gy': 2.1},   # right central incisor
        '21': {'x':  0.45, 'w': 0.85, 'h': 1.8, 'gy': 2.1},   # left central incisor
        '22': {'x':  1.40, 'w': 0.7, 'h': 1.5, 'gy': 1.9},    # left lateral incisor
    }

    for t_label, info in tooth_info.items():
        val = tooth_vals.get(t_label, 0)
        color = cmap(norm(val))

        # Tooth body (rounded rectangle)
        x = info['x'] - info['w'] / 2
        y = info['gy'] - info['h']
        rect = FancyBboxPatch(
            (x, y), info['w'], info['h'],
            boxstyle="round,pad=0.05",
            facecolor=color, edgecolor='#444444', linewidth=0.8, zorder=2
        )
        ax.add_patch(rect)

        # Tooth label and value
        cx = info['x']
        cy = info['gy'] - info['h'] / 2
        ax.text(cx, cy + 0.15, t_label, ha='center', va='center',
                fontsize=6, fontweight='bold', color='#333333', zorder=3)
        ax.text(cx, cy - 0.2, f'{val:.1f}', ha='center', va='center',
                fontsize=5.5, color='#333333', zorder=3)

    # Title and n
    ax.text(0, 3.3, title, ha='center', va='center', fontsize=7.5, fontweight='bold')
    ax.text(0, -0.3, f'n = {n}', ha='center', va='center', fontsize=6, color='gray')


# =====================================================================
# CREATE FIGURE: 2 rows (airway) × 4 cols (laryngoscope), each with F+M
# =====================================================================
fig, axes = plt.subplots(2, 8, figsize=(18, 8))

plt.rcParams.update({
    'font.family': 'Arial',
    'font.size': 8,
})

for row_idx, (aw, aw_label) in enumerate([(1, 'Standard airway'), (2, 'Difficult airway')]):
    for col_idx, (lv, lname) in enumerate(laryng.items()):
        ax_f = axes[row_idx, col_idx * 2]
        ax_m = axes[row_idx, col_idx * 2 + 1]

        d_f = data[(aw, lv, 'female')]
        d_m = data[(aw, lv, 'male')]

        draw_upper_jaw(ax_f, d_f['vals'], title='Female', n=d_f['n'])
        draw_upper_jaw(ax_m, d_m['vals'], title='Male', n=d_m['n'])

        # Laryngoscope label above pair (only top row)
        if row_idx == 0:
            # Position between the two axes
            mid_x = (ax_f.get_position().x0 + ax_m.get_position().x1) / 2
            fig.text(mid_x, 0.94, lname, ha='center', va='center',
                     fontsize=11, fontweight='bold')

    # Row label
    ax_left = axes[row_idx, 0]
    pos = ax_left.get_position()
    fig.text(0.01, (pos.y0 + pos.y1) / 2, aw_label,
             ha='left', va='center', fontsize=11, fontweight='bold', rotation=90)

# Colorbar
fig.subplots_adjust(bottom=0.14, top=0.88, left=0.04, right=0.96, hspace=0.3, wspace=0.15)
cbar_ax = fig.add_axes([0.25, 0.04, 0.5, 0.025])
sm = cm.ScalarMappable(cmap=cmap, norm=norm)
sm.set_array([])
cbar = fig.colorbar(sm, cax=cbar_ax, orientation='horizontal')
cbar.set_label('Median force (N)', fontsize=9, fontweight='bold')
cbar.ax.tick_params(labelsize=8)

# Main title
fig.suptitle('Fig 5.  Dental force heatmap on upper incisors — SLT only',
             fontsize=13, fontweight='bold', y=0.98)

fig5_path = os.path.join(tempfile.gettempdir(), 'fig5_tooth_heatmap.png')
fig.savefig(fig5_path, dpi=300, bbox_inches='tight')
plt.close(fig)
print(f"\nFig 5 saved: {fig5_path}")

# =====================================================================
# ADD TO figures.docx or create standalone
# =====================================================================
doc = Document()
section = doc.sections[-1]
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(1.5)
section.right_margin = Cm(1.5)

# Landscape
from docx.enum.section import WD_ORIENT
section.orientation = WD_ORIENT.LANDSCAPE
new_w, new_h = section.page_height, section.page_width
section.page_width = new_w
section.page_height = new_h

# Figure label
p_label = doc.add_paragraph()
p_label.alignment = WD_ALIGN_PARAGRAPH.LEFT
run_label = p_label.add_run('Fig 5')
run_label.bold = True
run_label.font.size = Pt(11)
run_label.font.name = 'Arial'
p_label.paragraph_format.space_after = Pt(4)

# Image
p_img = doc.add_paragraph()
p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
run_img = p_img.add_run()
run_img.add_picture(fig5_path, width=Inches(9))
p_img.paragraph_format.space_after = Pt(6)

# Caption
p_cap = doc.add_paragraph()
p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
run_cap = p_cap.add_run(
    'Fig 5. Dental force heatmap on upper incisors by laryngoscope type, airway difficulty, '
    'and sex (single-lumen tube only). '
    'Each diagram shows a simplified frontal view of the upper four incisors '
    '(teeth 12, 11, 21, 22). Tooth colour intensity represents the median maximum force (N) '
    'exerted during intubation, with darker shades indicating higher forces (see colour scale). '
    'Columns represent laryngoscope types (Macintosh, GlideScope, King Vision, C-MAC); '
    'rows represent airway difficulty (standard vs. difficult). '
    'Within each panel, the left jaw shows female and the right jaw shows male anaesthesiologists. '
    'Matched pairs with education levels kept separate (residents/specialists/attendings). '
    f'Total matched pairs (SLT): {len(df_w)}.'
)
run_cap.font.size = Pt(8)
run_cap.font.name = 'Arial'
run_cap.italic = True

doc.save('figures_heatmap.docx')
print("Saved: figures_heatmap.docx")
