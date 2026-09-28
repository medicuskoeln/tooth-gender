import pandas as pd
import numpy as np
from scipy.stats import mannwhitneyu, chi2_contingency, fisher_exact, shapiro, ttest_ind, kruskal, f_oneway
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

# --- Daten einlesen und auf Probandenebene reduzieren ---
df = pd.read_excel('roh1.xlsx')

proband_cols = ['PROBAND', 'Geschlecht_w1_m2', 'Ausbild_WA1_FA2_OA3',
                'Anz_SLT_4', 'Anz_DLT_4', 'Anz_MACI_4', 'Anz_GLIDE_CMAC_4', 'Anz_KingV_4',
                'NL_SLT_Easy_5', 'Nl_SLT_Diff_6', 'Nl_DLT_Easy_7', 'NI_DLT_Diff_8',
                'DA_SLT_Easy_9', 'DA_SLT_Diff_10', 'DA_DLT_Easy_11', 'DA_DLT_Diff_12',
                'Dent_Risk_Low_13', 'Dent_Risk_High_14', 'Dent_Damage_SLT_15', 'Dent_Damage_DLT_16']

df_prob = df[proband_cols].drop_duplicates(subset='PROBAND')

# 3 Gruppen nach Ausbildung
wa = df_prob[df_prob['Ausbild_WA1_FA2_OA3'] == 1]  # Assistenzärzte
fa = df_prob[df_prob['Ausbild_WA1_FA2_OA3'] == 2]  # Fachärzte
oa = df_prob[df_prob['Ausbild_WA1_FA2_OA3'] == 3]  # Oberärzte

n_wa, n_fa, n_oa = len(wa), len(fa), len(oa)
print(f"Probanden gesamt: {len(df_prob)}")
print(f"  Assistenzärzte (WA): {n_wa}")
print(f"  Fachärzte (FA): {n_fa}")
print(f"  Oberärzte (OA): {n_oa}")

# --- Geschlecht pro Gruppe ---
def geschlecht_info(grp, n_grp):
    n_w = (grp['Geschlecht_w1_m2'] == 1).sum()
    n_m = (grp['Geschlecht_w1_m2'] == 2).sum()
    return n_w, n_m

# --- Kontinuierliche Variable: 3-Gruppen-Test ---
def test_continuous_3(col, g1, g2, g3):
    """Kruskal-Wallis oder ANOVA für 3 Gruppen."""
    v1 = pd.to_numeric(g1[col], errors='coerce').dropna()
    v2 = pd.to_numeric(g2[col], errors='coerce').dropna()
    v3 = pd.to_numeric(g3[col], errors='coerce').dropna()
    if len(v1) < 3 or len(v2) < 3 or len(v3) < 3:
        return None, None, None, None, None

    # Normalverteilung prüfen
    try:
        p1 = shapiro(v1)[1]
        p2 = shapiro(v2)[1]
        p3 = shapiro(v3)[1]
    except:
        p1, p2, p3 = 0, 0, 0

    if p1 > 0.05 and p2 > 0.05 and p3 > 0.05:
        _, p_val = f_oneway(v1, v2, v3)
        test_name = "ANOVA"
        def fmt(v): return f"{v.mean():.1f} ± {v.std():.1f}"
    else:
        _, p_val = kruskal(v1, v2, v3)
        test_name = "Kruskal-Wallis"
        def fmt(v):
            iqr = v.quantile(0.75) - v.quantile(0.25)
            return f"{v.median():.1f} ({iqr:.1f})"

    return fmt(v1), fmt(v2), fmt(v3), p_val, test_name

# --- Kategorische Variable: 3-Gruppen ---
def test_cat_3(col, g1, g2, g3, labels_map):
    """Chi² Test über 3 Gruppen für kategorische Variable."""
    rows = []
    all_vals = pd.concat([g1[col], g2[col], g3[col]]).dropna()
    categories = sorted(all_vals.unique())

    ct_array = []
    for cat in categories:
        ct_array.append([(g[col] == cat).sum() for g in [g1, g2, g3]])
    ct_array = np.array(ct_array)

    if ct_array.min() < 5 and len(categories) == 2:
        # Fisher nur für 2x2 möglich, bei 2x3 Chi² verwenden
        _, p_val, _, _ = chi2_contingency(ct_array)
        test_name = "Chi²"
    else:
        _, p_val, _, _ = chi2_contingency(ct_array)
        test_name = "Chi²"

    for cat in categories:
        label = labels_map.get(cat, str(cat))
        n1 = (g1[col] == cat).sum()
        n2 = (g2[col] == cat).sum()
        n3 = (g3[col] == cat).sum()
        rows.append({
            'Variable': f"  {label}, n (%)",
            'WA': f"{n1} ({100*n1/len(g1):.1f}%)" if len(g1) > 0 else "-",
            'FA': f"{n2} ({100*n2/len(g2):.1f}%)" if len(g2) > 0 else "-",
            'OA': f"{n3} ({100*n3/len(g3):.1f}%)" if len(g3) > 0 else "-",
        })
    return rows, p_val, test_name

# --- Ergebnisse sammeln ---
results = []

# 1) Geschlecht pro Gruppe (kategorisch)
geschl_rows, p_geschl, t_geschl = test_cat_3(
    'Geschlecht_w1_m2', wa, fa, oa,
    {1: 'Weiblich', 2: 'Männlich'}
)
results.append({'header': 'Geschlecht', 'rows': geschl_rows, 'p': p_geschl, 'test': t_geschl})

# 2) Kontinuierliche Variablen
cont_vars = [
    ('Anz_SLT_4', 'Anzahl SLT'),
    ('Anz_DLT_4', 'Anzahl DLT'),
    ('Anz_MACI_4', 'Anzahl Mac'),
    ('Anz_GLIDE_CMAC_4', 'Anzahl Glide/CMAC'),
    ('Anz_KingV_4', 'Anzahl KingVision'),
    ('NL_SLT_Easy_5', 'NL SLT Easy'),
    ('Nl_SLT_Diff_6', 'NL SLT Diff'),
    ('Nl_DLT_Easy_7', 'NL DLT Easy'),
    ('NI_DLT_Diff_8', 'NI DLT Diff'),
    ('DA_SLT_Easy_9', 'DA SLT Easy'),
    ('DA_SLT_Diff_10', 'DA SLT Diff'),
    ('DA_DLT_Easy_11', 'DA DLT Easy'),
    ('DA_DLT_Diff_12', 'DA DLT Diff'),
    ('Dent_Risk_Low_13', 'Dent Risk Low'),
    ('Dent_Risk_High_14', 'Dent Risk High'),
    ('Dent_Damage_SLT_15', 'Dent Damage SLT'),
    ('Dent_Damage_DLT_16', 'Dent Damage DLT'),
]

cont_results = []
for col, label in cont_vars:
    r = test_continuous_3(col, wa, fa, oa)
    if r[0] is None:
        continue
    s_wa, s_fa, s_oa, p_val, test_name = r
    cont_results.append({
        'Variable': label,
        'WA': s_wa,
        'FA': s_fa,
        'OA': s_oa,
        'p': f"{p_val:.4f}",
        'test': test_name
    })

# --- Word-Dokument erstellen ---
doc = Document()

section = doc.sections[-1]
section.orientation = WD_ORIENT.LANDSCAPE
new_width, new_height = section.page_height, section.page_width
section.page_width = new_width
section.page_height = new_height
section.left_margin = Cm(1.5)
section.right_margin = Cm(1.5)

doc.add_heading('Tabelle 1: Vergleich Assistenzärzte vs. Fachärzte vs. Oberärzte', level=1)
doc.add_paragraph(
    f"Probanden gesamt: {len(df_prob)} "
    f"(Assistenzärzte: {n_wa}, Fachärzte: {n_fa}, Oberärzte: {n_oa})\n"
    "Kontinuierliche Variablen: Mittelwert ± SD (ANOVA) oder Median (IQR) (Kruskal-Wallis).\n"
    "Kategorische Variablen: n (%) (Chi²-Test)."
)

# Tabelle: Variable | WA | FA | OA | p-Wert | Test
table = doc.add_table(rows=1, cols=6)
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# Header
hdr = table.rows[0].cells
headers = ['Variable',
           f'Assistenzärzte\n(n={n_wa})',
           f'Fachärzte\n(n={n_fa})',
           f'Oberärzte\n(n={n_oa})',
           'p-Wert', 'Test']
for i, txt in enumerate(headers):
    hdr[i].text = txt
    for p in hdr[i].paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(9)
            run.font.name = 'Arial'

# Anaesthesiology-Stil: nur horizontale Linien
def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{kwargs.get("top", "nil")}" w:sz="{kwargs.get("top_sz", "0")}" w:space="0" w:color="{kwargs.get("top_color", "000000")}"/>'
        f'<w:bottom w:val="{kwargs.get("bottom", "nil")}" w:sz="{kwargs.get("bottom_sz", "0")}" w:space="0" w:color="{kwargs.get("bottom_color", "000000")}"/>'
        f'<w:left w:val="nil" w:sz="0" w:space="0" w:color="000000"/>'
        f'<w:right w:val="nil" w:sz="0" w:space="0" w:color="000000"/>'
        f'</w:tcBorders>')
    tcPr.append(tcBorders)

table.style = 'Normal Table'

# Obere Linie
for cell in table.rows[0].cells:
    set_cell_border(cell, top="single", top_sz="12", top_color="000000",
                    bottom="single", bottom_sz="8", bottom_color="000000")

# Kategorische Variablen (Geschlecht)
for block in results:
    # Header-Zeile
    row = table.add_row().cells
    row[0].text = block['header']
    row[4].text = f"{block['p']:.4f}"
    row[5].text = block['test']
    for p_elem in row[0].paragraphs:
        for run in p_elem.runs:
            run.bold = True
            run.font.size = Pt(9)
            run.font.name = 'Arial'
    for i in [4, 5]:
        for p_elem in row[i].paragraphs:
            p_elem.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p_elem.runs:
                run.font.size = Pt(9)
                run.font.name = 'Arial'
    # Detail-Zeilen
    for sub in block['rows']:
        row = table.add_row().cells
        row[0].text = sub['Variable']
        row[1].text = sub['WA']
        row[2].text = sub['FA']
        row[3].text = sub['OA']
        for i in range(6):
            for p_elem in row[i].paragraphs:
                p_elem.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
                for run in p_elem.runs:
                    run.font.size = Pt(9)
                    run.font.name = 'Arial'

# Kontinuierliche Variablen
for r in cont_results:
    row = table.add_row().cells
    row[0].text = r['Variable']
    row[1].text = r['WA']
    row[2].text = r['FA']
    row[3].text = r['OA']
    row[4].text = r['p']
    row[5].text = r['test']
    for i in range(6):
        for p_elem in row[i].paragraphs:
            p_elem.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p_elem.runs:
                run.font.size = Pt(9)
                run.font.name = 'Arial'
                # p-Wert fett wenn signifikant
                if i == 4:
                    try:
                        if float(r['p']) < 0.05:
                            run.bold = True
                    except:
                        pass

# Untere Linie
for cell in table.rows[-1].cells:
    set_cell_border(cell, bottom="single", bottom_sz="12", bottom_color="000000")

# Zellenabstand
for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            pf = p.paragraph_format
            pf.space_before = Pt(1)
            pf.space_after = Pt(1)

doc.save('Tab1-oa.docx')
print("\nTabelle gespeichert in Tab1-oa.docx")
