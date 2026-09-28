import pandas as pd
import numpy as np
from scipy.stats import mannwhitneyu, chi2_contingency, fisher_exact, shapiro, ttest_ind
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

# --- Data ---
df = pd.read_excel('roh1.xlsx')
proband_cols = ['PROBAND', 'Geschlecht_w1_m2', 'Ausbild_WA1_FA2_OA3',
                'Anz_SLT_4', 'Anz_DLT_4', 'Anz_MACI_4', 'Anz_GLIDE_CMAC_4', 'Anz_KingV_4',
                'NL_SLT_Easy_5', 'Nl_SLT_Diff_6', 'Nl_DLT_Easy_7', 'NI_DLT_Diff_8',
                'DA_SLT_Easy_9', 'DA_SLT_Diff_10', 'DA_DLT_Easy_11', 'DA_DLT_Diff_12',
                'Dent_Risk_Low_13', 'Dent_Risk_High_14', 'Dent_Damage_SLT_15', 'Dent_Damage_DLT_16']
df_prob = df[proband_cols].drop_duplicates(subset='PROBAND')
df_prob = df_prob[df_prob['Geschlecht_w1_m2'].isin([1, 2])]

female = df_prob[df_prob['Geschlecht_w1_m2'] == 1]
male = df_prob[df_prob['Geschlecht_w1_m2'] == 2]
nf, nm = len(female), len(male)
print(f"Total: n={len(df_prob)} (female={nf}, male={nm})")

# --- Helpers ---
def fmt_p(p_val):
    return f"{p_val:.3f}" if p_val >= 0.001 else "<0.001"

def test_continuous(col):
    vf = pd.to_numeric(female[col], errors='coerce').dropna()
    vm = pd.to_numeric(male[col], errors='coerce').dropna()
    if len(vf) < 3 or len(vm) < 3:
        return '-', '-', '-'
    try:
        ps_f, ps_m = shapiro(vf)[1], shapiro(vm)[1]
    except:
        ps_f, ps_m = 0, 0
    if ps_f > 0.05 and ps_m > 0.05:
        _, p_val = ttest_ind(vf, vm)
        sf = f"{vf.mean():.1f} \u00b1 {vf.std():.1f}"
        sm = f"{vm.mean():.1f} \u00b1 {vm.std():.1f}"
    else:
        _, p_val = mannwhitneyu(vf, vm, alternative='two-sided')
        def fmt(v):
            iqr = v.quantile(0.75) - v.quantile(0.25)
            return f"{v.median():.1f} ({iqr:.1f})"
        sf, sm = fmt(vf), fmt(vm)
    return sf, sm, fmt_p(p_val)

# --- Build rows ---
rows = []

# Training level (categorical: 3 levels)
ausbild_map = {1: 'Residents', 2: 'Specialists', 3: 'Attendings'}
ct_data = []
for val, label in ausbild_map.items():
    nf_cat = (female['Ausbild_WA1_FA2_OA3'] == val).sum()
    nm_cat = (male['Ausbild_WA1_FA2_OA3'] == val).sum()
    ct_data.append([nf_cat, nm_cat])
ct_array = np.array(ct_data)
_, p_cat, _, _ = chi2_contingency(ct_array)

for i, (val, label) in enumerate(ausbild_map.items()):
    nf_cat = ct_data[i][0]
    nm_cat = ct_data[i][1]
    p_str = fmt_p(p_cat) if i == 0 else ''
    rows.append((
        f'  {label}, n (%)',
        f"{nf_cat} ({100*nf_cat/nf:.0f})",
        f"{nm_cat} ({100*nm_cat/nm:.0f})",
        p_str
    ))

# Continuous variables
cont_vars = [
    ('Anz_SLT_4', 'No. of SLT intubations'),
    ('Anz_DLT_4', 'No. of DLT intubations'),
    ('Anz_MACI_4', 'No. of Macintosh uses'),
    ('Anz_GLIDE_CMAC_4', 'No. of GlideScope/C-MAC uses'),
    ('Anz_KingV_4', 'No. of King Vision uses'),
    ('NL_SLT_Easy_5', 'NL SLT easy'),
    ('Nl_SLT_Diff_6', 'NL SLT difficult'),
    ('Nl_DLT_Easy_7', 'NL DLT easy'),
    ('NI_DLT_Diff_8', 'NL DLT difficult'),
    ('DA_SLT_Easy_9', 'DA SLT easy'),
    ('DA_SLT_Diff_10', 'DA SLT difficult'),
    ('DA_DLT_Easy_11', 'DA DLT easy'),
    ('DA_DLT_Diff_12', 'DA DLT difficult'),
    ('Dent_Risk_Low_13', 'Dental risk low'),
    ('Dent_Risk_High_14', 'Dental risk high'),
    ('Dent_Damage_SLT_15', 'Dental damage SLT'),
    ('Dent_Damage_DLT_16', 'Dental damage DLT'),
]

for col, label in cont_vars:
    sf, sm, p_str = test_continuous(col)
    rows.append((label, sf, sm, p_str))

# --- Word Document ---
doc = Document()

section = doc.sections[-1]
section.orientation = WD_ORIENT.LANDSCAPE
w, h = section.page_height, section.page_width
section.page_width = w
section.page_height = h
section.left_margin = Cm(2)
section.right_margin = Cm(2)
section.top_margin = Cm(1.5)

title = doc.add_heading('', level=1)
run = title.add_run(
    'Table 1. Demographic and Experience Characteristics of Participants by Sex'
)
run.font.size = Pt(11)
run.font.name = 'Arial'
run.bold = True

legend = doc.add_paragraph()
legend.paragraph_format.space_after = Pt(4)
run = legend.add_run(
    'Data are presented as mean \u00b1 SD or median (IQR) depending on distribution '
    '(Shapiro\u2013Wilk test). P values from independent-samples t-test or Mann\u2013Whitney U test '
    'for continuous variables and \u03c7\u00b2 test for categorical variables. '
    'SLT = single-lumen tube; DLT = double-lumen tube; NL = normal laryngoscopy; '
    'DA = difficult airway; C-MAC = C-MAC videolaryngoscope.'
)
run.font.size = Pt(8)
run.font.name = 'Arial'
run.italic = True

# Table: Variable | Female | Male | P
table = doc.add_table(rows=1, cols=4)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Normal Table'

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{kwargs.get("top","nil")}" w:sz="{kwargs.get("top_sz","0")}" w:space="0" w:color="{kwargs.get("top_color","000000")}"/>'
        f'<w:bottom w:val="{kwargs.get("bottom","nil")}" w:sz="{kwargs.get("bottom_sz","0")}" w:space="0" w:color="{kwargs.get("bottom_color","000000")}"/>'
        f'<w:left w:val="nil" w:sz="0" w:space="0" w:color="000000"/>'
        f'<w:right w:val="nil" w:sz="0" w:space="0" w:color="000000"/>'
        f'</w:tcBorders>')
    tcPr.append(tcBorders)

def style_cell(cell, text, bold=False, size=9, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = text
    for p in cell.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.bold = bold
            r.font.size = Pt(size)
            r.font.name = 'Arial'

# Header
hdr = table.rows[0]
style_cell(hdr.cells[0], 'Variable', bold=True, align=WD_ALIGN_PARAGRAPH.LEFT)
style_cell(hdr.cells[1], f'Female (n = {nf})', bold=True)
style_cell(hdr.cells[2], f'Male (n = {nm})', bold=True)
style_cell(hdr.cells[3], 'P value', bold=True)

for cell in hdr.cells:
    set_cell_border(cell, top="single", top_sz="12", top_color="000000",
                    bottom="single", bottom_sz="8", bottom_color="000000")

# "Training level" label row
tr_row = table.add_row()
style_cell(tr_row.cells[0], 'Training level, n (%)', bold=False, align=WD_ALIGN_PARAGRAPH.LEFT)

# Data rows
for label, val_f, val_m, p_str in rows:
    r = table.add_row()
    style_cell(r.cells[0], label, align=WD_ALIGN_PARAGRAPH.LEFT)
    style_cell(r.cells[1], val_f)
    style_cell(r.cells[2], val_m)
    style_cell(r.cells[3], p_str)
    # Bold significant p
    if p_str and p_str != '-':
        try:
            pv = float(p_str) if p_str != '<0.001' else 0.0001
            if pv < 0.05:
                for p_elem in r.cells[3].paragraphs:
                    for run in p_elem.runs:
                        run.bold = True
        except:
            pass

# Bottom border
for cell in table.rows[-1].cells:
    set_cell_border(cell, bottom="single", bottom_sz="12", bottom_color="000000")

# Compact spacing
for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            pf = p.paragraph_format
            pf.space_before = Pt(1)
            pf.space_after = Pt(1)

doc.save('Tab1-oa.docx')
print("Saved: Tab1-oa.docx")
