import pandas as pd
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
import matplotlib.path as mpath
import matplotlib.patches as mpatches
from matplotlib.colors import Normalize
from matplotlib import cm
from scipy.stats import wilcoxon, shapiro, ttest_rel
from docx import Document
from docx.shared import Pt, Cm, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_ORIENT
import os, tempfile

# =====================================================================
# DATA LOADING
# =====================================================================
df_all = pd.read_excel('roh1.xlsx')
df_all = df_all[df_all['Geschlecht_w1_m2'].isin([1, 2])]

# OA und FA zusammenfassen: OA (3) → FA (2) => FAO-Gruppe
df_all['Ausbild_WA1_FA2_OA3'] = df_all['Ausbild_WA1_FA2_OA3'].replace({3: 2})

laryng = {1: 'Macintosh', 2: 'GlideScope', 3: 'King Vision', 4: 'C-MAC'}
teeth_cols = [
    ('res_12_max', '12'),
    ('res_11_max', '11'),
    ('res_21_max', '21'),
    ('res_22_max', '22'),
]
matching_cols = ['Ausbild_WA1_FA2_OA3', 'Laryng_diff_1Mac_2Glide_3King_4CMAC', 'tubus', 'atemweg']

# =====================================================================
# FUNCTION: match & compute for a given tube filter
# =====================================================================
def compute_data(df_input):
    """Match pairs, compute medians and p-values. Returns data, sig_data, n_total."""
    matched_pairs = []
    for _, group in df_input.groupby(matching_cols):
        w = group[group['Geschlecht_w1_m2'] == 1]
        m = group[group['Geschlecht_w1_m2'] == 2]
        n = min(len(w), len(m))
        if n > 0:
            matched_pairs.append((
                w.sample(n=n, random_state=42).reset_index(drop=True),
                m.sample(n=n, random_state=42).reset_index(drop=True)
            ))
    d_w = pd.concat([w for w, _ in matched_pairs], ignore_index=True)
    d_m = pd.concat([m for _, m in matched_pairs], ignore_index=True)

    data, sig_data = {}, {}
    for aw in [1, 2]:
        for lv in laryng:
            sub_f = d_w[(d_w['atemweg'] == aw) & (d_w['Laryng_diff_1Mac_2Glide_3King_4CMAC'] == lv)]
            sub_m = d_m[(d_m['atemweg'] == aw) & (d_m['Laryng_diff_1Mac_2Glide_3King_4CMAC'] == lv)]
            vals_f, vals_m, pvals = {}, {}, {}
            for col, label in teeth_cols:
                vf = pd.to_numeric(sub_f[col], errors='coerce')
                vm = pd.to_numeric(sub_m[col], errors='coerce')
                mask = (~vf.isna()) & (~vm.isna())
                vf_c, vm_c = vf[mask], vm[mask]
                vals_f[label] = vf_c.median() if len(vf_c) > 0 else 0
                vals_m[label] = vm_c.median() if len(vm_c) > 0 else 0
                if len(vf_c) >= 3:
                    diff = vf_c.values - vm_c.values
                    try: p_shap = shapiro(diff)[1]
                    except: p_shap = 0
                    if p_shap > 0.05:
                        _, p_val = ttest_rel(vf_c, vm_c)
                    else:
                        try: _, p_val = wilcoxon(vf_c, vm_c)
                        except: p_val = 1.0
                else:
                    p_val = 1.0
                pvals[label] = p_val
            n_pairs = len(sub_f)
            data[(aw, lv, 'female')] = {'vals': vals_f, 'n': n_pairs}
            data[(aw, lv, 'male')]   = {'vals': vals_m, 'n': n_pairs}
            sig_data[(aw, lv)] = pvals
    return data, sig_data, len(d_w)

# Compute for SLT, DLT, and combined
datasets = {}
for tube_label, tube_filter in [('SLT', 1), ('DLT', 2), ('SLT + DLT', None)]:
    if tube_filter is not None:
        df_sub = df_all[df_all['tubus'] == tube_filter]
    else:
        df_sub = df_all.copy()
    d, s, n = compute_data(df_sub)
    datasets[tube_label] = {'data': d, 'sig': s, 'n': n}
    print(f"\n{tube_label}: {len(df_sub)} rows, {n} matched pairs")
    for aw, aw_l in [(1,'std'),(2,'dif')]:
        for lv in laryng:
            pv = s[(aw,lv)]
            sig_str = " ".join(f"{t}={'*' if pv[t]<0.05 else 'ns'}" for _,t in teeth_cols)
            print(f"  {aw_l} {laryng[lv]:12s} n={d[(aw,lv,'female')]['n']:3d}  {sig_str}")

# Global color scale across ALL datasets
all_medians = []
for ds in datasets.values():
    for v in ds['data'].values():
        all_medians.extend(v['vals'].values())
vmin = 0
vmax = max(all_medians) * 1.05
print(f"\nGlobal color scale: {vmin:.1f} - {vmax:.1f} N")

# =====================================================================
# STYLE
# =====================================================================
cmap = cm.YlOrRd
norm = Normalize(vmin=vmin, vmax=vmax)
plt.rcParams.update({'font.family': 'Arial', 'font.size': 8})

# =====================================================================
# REALISTIC TOOTH SHAPES (Bezier curves)
# =====================================================================
def tooth_central(cx, base_y, w, h):
    """Upper central incisor: roots UP into maxilla, crown DOWN."""
    hw = w / 2
    rw = hw * 0.42
    crown_h = h * 0.58
    root_h = h * 0.42
    cy = base_y  # cervical (gum) line

    P = mpath.Path
    v = [
        # Start at cervical left, go UP to root apex
        (cx - rw,          cy),
        (cx - rw * 1.08,   cy + root_h * 0.30),
        (cx - rw * 0.55,   cy + root_h * 0.90),
        (cx,               cy + root_h),           # apex (top)
        (cx + rw * 0.55,   cy + root_h * 0.90),
        (cx + rw * 1.08,   cy + root_h * 0.30),
        (cx + rw,          cy),                     # cervical right
        # Crown right side (going DOWN)
        (cx + hw * 0.75,   cy - crown_h * 0.12),
        (cx + hw * 0.92,   cy - crown_h * 0.35),
        (cx + hw,          cy - crown_h * 0.55),    # max width
        # Incisal edge right
        (cx + hw * 0.98,   cy - crown_h * 0.82),
        (cx + hw * 0.88,   cy - crown_h * 0.97),
        (cx + hw * 0.35,   cy - crown_h - 0.015),  # right mamellon
        # Incisal middle
        (cx + hw * 0.12,   cy - crown_h + 0.02),
        (cx - hw * 0.12,   cy - crown_h + 0.02),
        (cx - hw * 0.35,   cy - crown_h - 0.015),  # left mamellon
        # Incisal left
        (cx - hw * 0.88,   cy - crown_h * 0.97),
        (cx - hw * 0.98,   cy - crown_h * 0.82),
        (cx - hw,          cy - crown_h * 0.55),    # max width left
        # Crown left side
        (cx - hw * 0.92,   cy - crown_h * 0.35),
        (cx - hw * 0.75,   cy - crown_h * 0.12),
        (cx - rw,          cy),                     # close
    ]
    c = [
        P.MOVETO,
        P.CURVE4, P.CURVE4, P.CURVE4,  # root up left
        P.CURVE4, P.CURVE4, P.CURVE4,  # root up right
        P.CURVE4, P.CURVE4, P.CURVE4,  # crown right
        P.CURVE4, P.CURVE4, P.CURVE4,  # incisal right
        P.CURVE4, P.CURVE4, P.CURVE4,  # incisal middle
        P.CURVE4, P.CURVE4, P.CURVE4,  # incisal left
        P.CURVE4, P.CURVE4, P.CURVE4,  # crown left
    ]
    return P(v, c)


def tooth_lateral(cx, base_y, w, h):
    """Upper lateral incisor: roots UP, crown DOWN."""
    hw = w / 2
    rw = hw * 0.38
    crown_h = h * 0.54
    root_h = h * 0.46

    cy = base_y

    P = mpath.Path
    v = [
        # Root goes UP
        (cx - rw,          cy),
        (cx - rw * 1.05,   cy + root_h * 0.28),
        (cx - rw * 0.48,   cy + root_h * 0.88),
        (cx,               cy + root_h),
        (cx + rw * 0.48,   cy + root_h * 0.88),
        (cx + rw * 1.05,   cy + root_h * 0.28),
        (cx + rw,          cy),
        # Crown right (going DOWN)
        (cx + hw * 0.68,   cy - crown_h * 0.12),
        (cx + hw * 0.88,   cy - crown_h * 0.32),
        (cx + hw,          cy - crown_h * 0.52),
        # Incisal (rounded, bottom)
        (cx + hw * 0.96,   cy - crown_h * 0.80),
        (cx + hw * 0.72,   cy - crown_h * 0.97),
        (cx,               cy - crown_h),
        (cx - hw * 0.72,   cy - crown_h * 0.97),
        (cx - hw * 0.96,   cy - crown_h * 0.80),
        (cx - hw,          cy - crown_h * 0.52),
        # Crown left
        (cx - hw * 0.88,   cy - crown_h * 0.32),
        (cx - hw * 0.68,   cy - crown_h * 0.12),
        (cx - rw,          cy),
    ]
    c = [
        P.MOVETO,
        P.CURVE4, P.CURVE4, P.CURVE4,
        P.CURVE4, P.CURVE4, P.CURVE4,
        P.CURVE4, P.CURVE4, P.CURVE4,
        P.CURVE4, P.CURVE4, P.CURVE4,
        P.CURVE4, P.CURVE4, P.CURVE4,
        P.CURVE4, P.CURVE4, P.CURVE4,
    ]
    return P(v, c)


# =====================================================================
# DRAW JAW
# =====================================================================
def draw_upper_jaw(ax, tooth_vals, sig_markers=None, title='', n=0):
    ax.set_xlim(-3.2, 3.2)
    ax.set_ylim(-2.8, 3.8)
    ax.set_aspect('equal')
    ax.axis('off')

    # Gingiva background at TOP (above roots, which point upward)
    gum_t = np.linspace(0.05 * np.pi, 0.95 * np.pi, 300)
    gum_x = 3.5 * np.cos(gum_t)
    gum_y = 2.2 * np.sin(gum_t) + 1.2  # arched above roots
    ax.fill_between(gum_x, gum_y, 5.0, color='#E8A0A0', alpha=0.3, zorder=0)

    # Festooned gum margin (scalloped, above cervical lines)
    mgn_x = np.linspace(-2.5, 2.5, 500)
    mgn_y = np.ones_like(mgn_x) * 0.55
    # Arc over each tooth (dips downward toward cervical line)
    for tcx, tw in [(-1.55, 0.82), (-0.52, 1.0), (0.52, 1.0), (1.55, 0.82)]:
        hw = tw / 2 * 0.8
        m = (mgn_x > tcx - hw) & (mgn_x < tcx + hw)
        local = (mgn_x[m] - tcx) / hw
        mgn_y[m] = 0.55 - 0.22 * (1 - local**2)
    # Papillae (pointing downward between teeth)
    for px in [-1.02, 0.0, 1.02]:
        m = np.abs(mgn_x - px) < 0.18
        local = (mgn_x[m] - px) / 0.18
        mgn_y[m] -= 0.35 * np.cos(local * np.pi / 2)

    ax.fill_between(mgn_x, mgn_y, 5.0, color='#D88E8E', alpha=0.6, zorder=4)
    ax.plot(mgn_x, mgn_y, color='#B06060', lw=0.9, zorder=5)

    # Tooth definitions: base_y is cervical line, roots go UP (+), crown goes DOWN (-)
    tdefs = {
        '12': {'cx': -1.55, 'w': 0.82, 'h': 2.6, 'base_y': -0.05, 'fn': tooth_lateral},
        '11': {'cx': -0.52, 'w': 1.0,  'h': 2.9, 'base_y':  0.10, 'fn': tooth_central},
        '21': {'cx':  0.52, 'w': 1.0,  'h': 2.9, 'base_y':  0.10, 'fn': tooth_central},
        '22': {'cx':  1.55, 'w': 0.82, 'h': 2.6, 'base_y': -0.05, 'fn': tooth_lateral},
    }

    for t_label, td in tdefs.items():
        val = tooth_vals.get(t_label, 0)
        color = cmap(norm(val))

        path = td['fn'](td['cx'], td['base_y'], td['w'], td['h'])

        # Tooth fill
        patch = mpatches.PathPatch(path, facecolor=color, edgecolor='#3a3a3a',
                                   lw=0.8, zorder=3, alpha=0.9)
        ax.add_patch(patch)

        # Light reflection (glossy highlight on crown)
        highlight_x = td['cx'] - td['w'] * 0.12
        crown_h = td['h'] * (0.58 if td['fn'] == tooth_central else 0.54)
        hl_y = td['base_y'] - crown_h * 0.45
        ax.plot([highlight_x, highlight_x],
                [hl_y - crown_h * 0.2, hl_y + crown_h * 0.2],
                color='white', lw=1.5, alpha=0.3, zorder=4, solid_capstyle='round')

        # Cervical line
        hw = td['w'] / 2
        cej_xs = np.linspace(td['cx'] - hw * 0.6, td['cx'] + hw * 0.6, 40)
        cej_ys = td['base_y'] - 0.06 * np.sin(np.pi * (cej_xs - (td['cx'] - hw * 0.6)) / (hw * 1.2))
        ax.plot(cej_xs, cej_ys, color='#777777', lw=0.35, zorder=4, alpha=0.5)

        # Labels (on the crown, which is now below cervical line)
        label_y = td['base_y'] - crown_h * 0.45
        ax.text(td['cx'], label_y + 0.18, t_label, ha='center', va='center',
                fontsize=8, fontweight='bold', color='#1a1a1a', zorder=6)
        ax.text(td['cx'], label_y - 0.35, f'{val:.1f} N', ha='center', va='center',
                fontsize=7, color='#222222', zorder=6)

        # Significance asterisk below incisal edge (at bottom of crown)
        if sig_markers and sig_markers.get(t_label, False):
            star_y = td['base_y'] - crown_h - 0.15
            ax.text(td['cx'], star_y, '*', ha='center', va='top',
                    fontsize=16, color='#CC0000', fontweight='bold', zorder=7)

    # Title above jaw
    ax.text(0, 3.5, title, ha='center', va='center', fontsize=10, fontweight='bold')
    ax.text(0, -2.5, f'n = {n}', ha='center', va='center', fontsize=8, color='#555555')


# =====================================================================
# GENERATE FIGURE for a dataset
# =====================================================================
def make_figure(data, sig_data, n_total, tube_label, fig_num):
    fig, axes = plt.subplots(4, 4, figsize=(14, 18))

    col_configs = [
        (1, 'female', 'Female\nStandard'),
        (1, 'male',   'Male\nStandard'),
        (2, 'female', 'Female\nDifficult'),
        (2, 'male',   'Male\nDifficult'),
    ]

    for row_idx, (lv, lname) in enumerate(laryng.items()):
        for col_idx, (aw, sex, col_title) in enumerate(col_configs):
            ax = axes[row_idx, col_idx]
            d = data[(aw, lv, sex)]
            pvals = sig_data[(aw, lv)]
            sig_marks = {t: pvals[t] < 0.05 for _, t in teeth_cols}

            draw_upper_jaw(ax, d['vals'], sig_markers=sig_marks,
                           title=col_title if row_idx == 0 else
                                 ('Female' if sex == 'female' else 'Male'),
                           n=d['n'])

    fig.subplots_adjust(bottom=0.07, top=0.90, left=0.08, right=0.96, hspace=0.22, wspace=0.12)

    # Column group headers
    ax_s1 = axes[0, 0].get_position()
    ax_s2 = axes[0, 1].get_position()
    ax_d1 = axes[0, 2].get_position()
    ax_d2 = axes[0, 3].get_position()
    fig.text((ax_s1.x0 + ax_s2.x1) / 2, 0.935, 'Standard Airway',
             ha='center', fontsize=14, fontweight='bold')
    fig.text((ax_d1.x0 + ax_d2.x1) / 2, 0.935, 'Difficult Airway',
             ha='center', fontsize=14, fontweight='bold')

    # Row labels
    for row_idx, (lv, lname) in enumerate(laryng.items()):
        pos = axes[row_idx, 0].get_position()
        fig.text(0.025, (pos.y0 + pos.y1) / 2, lname,
                 ha='left', va='center', fontsize=13, fontweight='bold', rotation=90)

    # Colorbar
    cbar_ax = fig.add_axes([0.25, 0.02, 0.5, 0.012])
    sm = cm.ScalarMappable(cmap=cmap, norm=norm)
    sm.set_array([])
    cbar = fig.colorbar(sm, cax=cbar_ax, orientation='horizontal')
    cbar.set_label('Median force (N)', fontsize=11, fontweight='bold')
    cbar.ax.tick_params(labelsize=9)

    fig.text(0.80, 0.018, '* P < 0.05 (paired test, female vs. male)',
             fontsize=10, color='#CC0000', fontstyle='italic', ha='left')

    fig.suptitle(f'Fig {fig_num}.  Dental force heatmap on upper incisors \u2014 {tube_label}',
                 fontsize=15, fontweight='bold', y=0.96)

    path = os.path.join(tempfile.gettempdir(), f'fig{fig_num}_heatmap_{tube_label.replace(" ","").replace("+","")}.png')
    fig.savefig(path, dpi=300, bbox_inches='tight')
    plt.close(fig)
    print(f"Fig {fig_num} ({tube_label}) saved")
    return path

# Generate all 3 figures
fig_paths = {}
for i, (tube_label, ds) in enumerate(datasets.items(), start=5):
    fig_paths[tube_label] = make_figure(
        ds['data'], ds['sig'], ds['n'], tube_label, fig_num=i
    )

# =====================================================================
# WORD DOCUMENT with all 3 figures
# =====================================================================
doc = Document()
section = doc.sections[-1]
section.orientation = WD_ORIENT.LANDSCAPE
new_w, new_h = section.page_height, section.page_width
section.page_width = new_w
section.page_height = new_h
section.top_margin = Cm(1.5)
section.bottom_margin = Cm(1.5)
section.left_margin = Cm(1.5)
section.right_margin = Cm(1.5)

def add_fig_page(doc, img_path, fig_num, tube_label, n_total, is_first=False):
    if not is_first:
        doc.add_page_break()

    p_label = doc.add_paragraph()
    p_label.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run_label = p_label.add_run(f'Fig {fig_num}')
    run_label.bold = True
    run_label.font.size = Pt(11)
    run_label.font.name = 'Arial'
    p_label.paragraph_format.space_after = Pt(4)

    p_img = doc.add_paragraph()
    p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_img = p_img.add_run()
    run_img.add_picture(img_path, width=Inches(9))
    p_img.paragraph_format.space_after = Pt(6)

    p_cap = doc.add_paragraph()
    p_cap.alignment = WD_ALIGN_PARAGRAPH.LEFT
    if tube_label == 'SLT + DLT':
        tube_desc = 'single-lumen and double-lumen tubes combined'
    elif tube_label == 'SLT':
        tube_desc = 'single-lumen tube only'
    else:
        tube_desc = 'double-lumen tube only'
    run_cap = p_cap.add_run(
        f'Fig {fig_num}. Dental force heatmap on upper incisors by laryngoscope type, '
        f'airway difficulty, and sex ({tube_desc}). '
        'Each panel shows a frontal view of the upper four incisors '
        '(teeth 12, 11, 21, 22) with tooth colour intensity representing the median maximum '
        'force (N) exerted during intubation (yellow = low force, red = high force; '
        'unified colour scale across all panels and figures). '
        'Rows represent laryngoscope types (Macintosh, GlideScope, King Vision, C-MAC). '
        'Within each airway condition, the left panel shows female, the right panel male '
        'anaesthesiologists. '
        'Red asterisks (*) below incisal edges indicate statistically significant paired '
        'female\u2013male differences (P < 0.05; Wilcoxon signed-rank test or paired t-test). '
        'Matched pairs with specialists and attendings combined (residents/FAO). '
        f'Total matched pairs: {n_total}.'
    )
    run_cap.font.size = Pt(8)
    run_cap.font.name = 'Arial'
    run_cap.italic = True

for i, (tube_label, ds) in enumerate(datasets.items(), start=5):
    add_fig_page(doc, fig_paths[tube_label], i, tube_label, ds['n'],
                 is_first=(i == 5))

doc.save('figures_heatmap.docx')
print("\nSaved: figures_heatmap.docx")
