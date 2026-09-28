import pandas as pd
import numpy as np
from scipy.stats import spearmanr
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

from stats_helpers import (
    bh_fdr, fmt_p, paired_test, mixed_effects_p, LEGEND_FDR_MEM,
)

# --- Daten einlesen ---
df = pd.read_excel('roh1.xlsx')
df = df[df['Geschlecht_w1_m2'].isin([1, 2])]
df['Ausbild_WA1_FA2_OA3'] = df['Ausbild_WA1_FA2_OA3'].replace({3: 2})

# --- Matching ---
matching_cols = ['Ausbild_WA1_FA2_OA3', 'Laryng_diff_1Mac_2Glide_3King_4CMAC', 'tubus', 'atemweg']

matched_pairs = []
for _, group in df.groupby(matching_cols):
    w = group[group['Geschlecht_w1_m2'] == 1]
    m = group[group['Geschlecht_w1_m2'] == 2]
    n = min(len(w), len(m))
    if n > 0:
        matched_pairs.append((
            w.sample(n=n, random_state=42).reset_index(drop=True),
            m.sample(n=n, random_state=42).reset_index(drop=True)
        ))

df_w = pd.concat([w for w, _ in matched_pairs], ignore_index=True)
df_m = pd.concat([m for _, m in matched_pairs], ignore_index=True)

print(f"Gematchte Paare gesamt: {len(df_w)}")

laryng_map = {1: 'Macintosh', 2: 'GlideScope', 3: 'King Vision', 4: 'C-MAC'}
airway_map = {1: 'Easy airway', 2: 'Difficult airway'}
outcome_vars = [
    ('Empf_Schwierig', 'Difficulty (0\u201310)'),
    ('Sicht_Cormack', 'C/L grade'),
    ('res_max_alle_zaehne', 'Total dental force (N)'),
]


def analyse_matched(col, dw, dm):
    vw = pd.to_numeric(dw[col], errors='coerce')
    vm = pd.to_numeric(dm[col], errors='coerce')
    mask = (~vw.isna()) & (~vm.isna())
    vw_v = vw[mask]
    vm_v = vm[mask]
    if len(vw_v) < 3:
        return '-', '-', np.nan, np.nan
    diff = vw_v.values - vm_v.values
    from scipy.stats import shapiro
    try:
        p_shap = shapiro(diff)[1]
    except Exception:
        p_shap = 0
    if p_shap > 0.05:
        sf = f"{vw_v.mean():.2f} \u00b1 {vw_v.std():.2f}"
        sm = f"{vm_v.mean():.2f} \u00b1 {vm_v.std():.2f}"
    else:
        def fmt(v):
            return f"{v.median():.2f} ({v.quantile(0.25):.2f}\u2013{v.quantile(0.75):.2f})"
        sf, sm = fmt(vw_v), fmt(vm_v)
    p_raw, _ = paired_test(vw_v.values, vm_v.values)
    p_mm = mixed_effects_p(dw.loc[mask.values], dm.loc[mask.values], col)
    return sf, sm, p_raw, p_mm


# --- Ergebnisse berechnen ---
table_data = []
all_p_raw = []
all_p_mm = []
for lv, lname in laryng_map.items():
    for aw, aw_label in airway_map.items():
        w_sub = df_w[(df_w['Laryng_diff_1Mac_2Glide_3King_4CMAC'] == lv) &
                     (df_w['atemweg'] == aw)].reset_index(drop=True)
        m_sub = df_m[(df_m['Laryng_diff_1Mac_2Glide_3King_4CMAC'] == lv) &
                     (df_m['atemweg'] == aw)].reset_index(drop=True)
        n_pairs = len(w_sub)
        for col, label in outcome_vars:
            sf, sm, p_raw, p_mm = analyse_matched(col, w_sub, m_sub)
            table_data.append({
                'laryng': lname,
                'laryng_code': lv,
                'airway': aw_label,
                'airway_code': aw,
                'var_label': label,
                'female': sf,
                'male': sm,
                'p_raw': p_raw,
                'p_mm': p_mm,
                'n_pairs': n_pairs,
            })
            all_p_raw.append(p_raw)
            all_p_mm.append(p_mm)

qfdr_all = bh_fdr(all_p_raw)
qmm_all = bh_fdr(all_p_mm)
for i, entry in enumerate(table_data):
    entry['q'] = qfdr_all[i]
    entry['qmm'] = qmm_all[i]

# Stichprobengrößen
n_info = {}
for lv, lname in laryng_map.items():
    for aw, aw_label in airway_map.items():
        nw = len(df_w[(df_w['Laryng_diff_1Mac_2Glide_3King_4CMAC'] == lv) & (df_w['atemweg'] == aw)])
        nm = len(df_m[(df_m['Laryng_diff_1Mac_2Glide_3King_4CMAC'] == lv) & (df_m['atemweg'] == aw)])
        n_info[(lv, aw)] = (nw, nm)
        print(f"{lname} \u2013 {aw_label}: {nw} Paare")

total_pairs = len(df_w)

# --- Word-Dokument ---
doc = Document()

section = doc.sections[-1]
section.orientation = WD_ORIENT.LANDSCAPE
w, h = section.page_height, section.page_width
section.page_width = w
section.page_height = h
section.left_margin = Cm(2)
section.right_margin = Cm(2)
section.top_margin = Cm(1.5)

title_p = doc.add_paragraph()
title_run = title_p.add_run('Table.')
title_run.bold = True
title_run.font.size = Pt(11)
title_run.font.name = 'Arial'
title_p.paragraph_format.space_after = Pt(2)

subtitle_p = doc.add_paragraph()
subtitle_run = subtitle_p.add_run(
    'Matched-Pair Comparison of Perceived Difficulty, Cormack\u2013Lehane Grade, and Total Dental Force '
    'by Laryngoscope Type and Airway Difficulty \u2014 Female vs. Male'
)
subtitle_run.font.size = Pt(9)
subtitle_run.font.name = 'Arial'
subtitle_p.paragraph_format.space_after = Pt(6)

ncols = 4
table = doc.add_table(rows=1, cols=ncols)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Normal Table'


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{kwargs.get("top","nil")}" w:sz="{kwargs.get("top_sz","0")}" w:space="0" w:color="{kwargs.get("top_color","000000")}"/>'
        f'<w:bottom w:val="{kwargs.get("bottom","nil")}" w:sz="{kwargs.get("bottom_sz","0")}" w:space="0" w:color="{kwargs.get("bottom_color","000000")}"/>'
        f'<w:left w:val="nil" w:sz="0" w:space="0" w:color="000000"/>'
        f'<w:right w:val="nil" w:sz="0" w:space="0" w:color="000000"/>'
        f'</w:tcBorders>')
    tcPr.append(tcBorders)


def style_cell(cell, text, bold=False, italic=False, size=9, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.italic = italic
    run.font.size = Pt(size)
    run.font.name = 'Arial'


def write_p_cell(cell, p_raw, p_fdr, p_mm):
    cell.text = ''
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    sig = ''
    try:
        if p_fdr is not None and not np.isnan(p_fdr) and p_fdr < 0.05:
            sig += '*'
    except Exception:
        pass
    try:
        if p_mm is not None and not np.isnan(p_mm) and p_mm < 0.05:
            sig += '\u2020'
    except Exception:
        pass
    r1 = para.add_run(fmt_p(p_raw) + sig)
    r1.font.size = Pt(8.5)
    r1.font.name = 'Arial'
    try:
        if p_raw is not None and not np.isnan(float(p_raw)) and float(p_raw) < 0.05:
            r1.bold = True
    except Exception:
        pass
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(f"qFDR={fmt_p(p_fdr)}")
    r2.font.size = Pt(7)
    r2.font.name = 'Arial'
    r2.italic = True
    p3 = cell.add_paragraph()
    p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p3.paragraph_format.space_before = Pt(0)
    p3.paragraph_format.space_after = Pt(0)
    r3 = p3.add_run(f"pMM={fmt_p(p_mm)}")
    r3.font.size = Pt(7)
    r3.font.name = 'Arial'
    r3.italic = True


# Header-Zeile
hdr = table.rows[0]
style_cell(hdr.cells[0], 'Variable', bold=True, size=9, align=WD_ALIGN_PARAGRAPH.LEFT)
style_cell(hdr.cells[1], f'Female (n = {len(df_w)})', bold=True, size=9)
style_cell(hdr.cells[2], f'Male (n = {len(df_m)})', bold=True, size=9)
style_cell(hdr.cells[3], 'P / qFDR / pMM', bold=True, size=9)

for cell in hdr.cells:
    set_cell_border(cell, top="single", top_sz="12", top_color="000000",
                    bottom="single", bottom_sz="8", bottom_color="000000")

prev_laryng = None
prev_airway = None

for entry in table_data:
    if entry['laryng'] != prev_laryng:
        lrow = table.add_row()
        n_easy = n_info[(entry['laryng_code'], 1)][0]
        n_diff = n_info[(entry['laryng_code'], 2)][0]
        style_cell(lrow.cells[0], f"{entry['laryng']} (n = {n_easy + n_diff} pairs)",
                   bold=True, size=9, align=WD_ALIGN_PARAGRAPH.LEFT)
        style_cell(lrow.cells[1], '', size=9)
        style_cell(lrow.cells[2], '', size=9)
        style_cell(lrow.cells[3], '', size=9)
        if prev_laryng is not None:
            for cell in lrow.cells:
                set_cell_border(cell, top="single", top_sz="4", top_color="000000")
        prev_laryng = entry['laryng']
        prev_airway = None

    if entry['airway'] != prev_airway:
        arow = table.add_row()
        n_aw = n_info[(entry['laryng_code'], entry['airway_code'])][0]
        style_cell(arow.cells[0], f"  {entry['airway']} (n = {n_aw})",
                   italic=True, size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT)
        style_cell(arow.cells[1], '', size=8.5)
        style_cell(arow.cells[2], '', size=8.5)
        style_cell(arow.cells[3], '', size=8.5)
        prev_airway = entry['airway']

    row = table.add_row()
    style_cell(row.cells[0], f"    {entry['var_label']}", size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT)
    style_cell(row.cells[1], entry['female'], size=8)
    style_cell(row.cells[2], entry['male'], size=8)
    write_p_cell(row.cells[3], entry['p_raw'], entry['q'], entry['p_mm'])

for cell in table.rows[-1].cells:
    set_cell_border(cell, bottom="single", bottom_sz="12", bottom_color="000000")

for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            pf = p.paragraph_format
            if pf.space_before is None:
                pf.space_before = Pt(1)
            if pf.space_after is None:
                pf.space_after = Pt(1)

# Legende (Tabelle 1)
legend = doc.add_paragraph()
legend.paragraph_format.space_before = Pt(6)
legend_text = (
    'Pairs matched on education level (residents vs. specialists/attendings combined), '
    'laryngoscope type, tube type, and airway difficulty. '
    'Data presented as mean \u00b1 SD or median (IQR) depending on distribution '
    '(Shapiro\u2013Wilk test on paired differences). '
    f'Total matched pairs: {total_pairs}. '
    'Difficulty rated on a scale of 0\u201310 (10 = most difficult). '
    'C/L = Cormack\u2013Lehane grade. '
    + LEGEND_FDR_MEM
)
run = legend.add_run(legend_text)
run.font.size = Pt(7.5)
run.font.name = 'Arial'
run.italic = True

# =============================================================
# Korrelationsanalyse: Schwierigkeit vs. C/L und vs. Druck
# =============================================================
def calc_corr(data, col_x, col_y):
    x = pd.to_numeric(data[col_x], errors='coerce')
    y = pd.to_numeric(data[col_y], errors='coerce')
    mask = (~x.isna()) & (~y.isna())
    x, y = x[mask], y[mask]
    if len(x) < 5:
        return '-', np.nan, 0
    rho, p = spearmanr(x, y)
    return f"{rho:.3f}", float(p), int(len(x))


corr_pairs = [
    ('Empf_Schwierig', 'Sicht_Cormack', 'Difficulty vs. C/L grade'),
    ('Empf_Schwierig', 'res_max_alle_zaehne', 'Difficulty vs. Total dental force'),
]

corr_results = []
all_corr_p_f = []
all_corr_p_m = []
for col_x, col_y, label in corr_pairs:
    for lv, lname in laryng_map.items():
        for aw, aw_label in airway_map.items():
            w_sub = df_w[(df_w['Laryng_diff_1Mac_2Glide_3King_4CMAC'] == lv) &
                         (df_w['atemweg'] == aw)]
            m_sub = df_m[(df_m['Laryng_diff_1Mac_2Glide_3King_4CMAC'] == lv) &
                         (df_m['atemweg'] == aw)]
            rho_f, p_f, n_f = calc_corr(w_sub, col_x, col_y)
            rho_m, p_m, n_m = calc_corr(m_sub, col_x, col_y)
            corr_results.append({
                'corr_label': label,
                'laryng': lname,
                'laryng_code': lv,
                'airway': aw_label,
                'airway_code': aw,
                'rho_f': rho_f, 'p_f': p_f, 'n_f': n_f,
                'rho_m': rho_m, 'p_m': p_m, 'n_m': n_m,
            })
            all_corr_p_f.append(p_f)
            all_corr_p_m.append(p_m)

q_corr_f = bh_fdr(all_corr_p_f)
q_corr_m = bh_fdr(all_corr_p_m)
for i, entry in enumerate(corr_results):
    entry['q_f'] = q_corr_f[i]
    entry['q_m'] = q_corr_m[i]

# --- Korrelationstabelle ---
doc.add_page_break()

corr_title_p = doc.add_paragraph()
corr_title_run = corr_title_p.add_run('Table.')
corr_title_run.bold = True
corr_title_run.font.size = Pt(11)
corr_title_run.font.name = 'Arial'
corr_title_p.paragraph_format.space_after = Pt(2)

corr_sub_p = doc.add_paragraph()
corr_sub_run = corr_sub_p.add_run(
    'Spearman Rank Correlation Between Perceived Difficulty and '
    'Cormack\u2013Lehane Grade or Total Dental Force \u2014 by Sex, Laryngoscope, '
    'and Airway Difficulty'
)
corr_sub_run.font.size = Pt(9)
corr_sub_run.font.name = 'Arial'
corr_sub_p.paragraph_format.space_after = Pt(6)

# 5 cols (Variable | Female ρ | Female P / qFDR | Male ρ | Male P / qFDR)
corr_table = doc.add_table(rows=1, cols=5)
corr_table.alignment = WD_TABLE_ALIGNMENT.CENTER
corr_table.style = 'Normal Table'

ch = corr_table.rows[0]
style_cell(ch.cells[0], 'Variable', bold=True, size=9, align=WD_ALIGN_PARAGRAPH.LEFT)
style_cell(ch.cells[1], 'Female \u03c1', bold=True, size=9)
style_cell(ch.cells[2], 'Female P / qFDR', bold=True, size=9)
style_cell(ch.cells[3], 'Male \u03c1', bold=True, size=9)
style_cell(ch.cells[4], 'Male P / qFDR', bold=True, size=9)

for cell in ch.cells:
    set_cell_border(cell, top="single", top_sz="12", top_color="000000",
                    bottom="single", bottom_sz="8", bottom_color="000000")


def write_corr_p_cell(cell, p_raw, p_fdr):
    cell.text = ''
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    sig = ''
    try:
        if p_fdr is not None and not np.isnan(p_fdr) and p_fdr < 0.05:
            sig += '*'
    except Exception:
        pass
    r1 = para.add_run(fmt_p(p_raw) + sig)
    r1.font.size = Pt(8)
    r1.font.name = 'Arial'
    try:
        if p_raw is not None and not np.isnan(float(p_raw)) and float(p_raw) < 0.05:
            r1.bold = True
    except Exception:
        pass
    p2 = cell.add_paragraph()
    p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p2.paragraph_format.space_before = Pt(0)
    p2.paragraph_format.space_after = Pt(0)
    r2 = p2.add_run(f"qFDR={fmt_p(p_fdr)}")
    r2.font.size = Pt(6.5)
    r2.font.name = 'Arial'
    r2.italic = True


prev_corr = None
prev_corr_laryng = None
for entry in corr_results:
    if entry['corr_label'] != prev_corr:
        crow = corr_table.add_row()
        style_cell(crow.cells[0], entry['corr_label'], bold=True, size=9,
                   align=WD_ALIGN_PARAGRAPH.LEFT)
        for c in crow.cells[1:]:
            style_cell(c, '', size=9)
        if prev_corr is not None:
            for c in crow.cells:
                set_cell_border(c, top="single", top_sz="8", top_color="000000")
        prev_corr = entry['corr_label']
        prev_corr_laryng = None

    if entry['laryng'] != prev_corr_laryng:
        lrow2 = corr_table.add_row()
        style_cell(lrow2.cells[0], f"  {entry['laryng']}", bold=True, italic=False,
                   size=8.5, align=WD_ALIGN_PARAGRAPH.LEFT)
        for c in lrow2.cells[1:]:
            style_cell(c, '', size=8.5)
        prev_corr_laryng = entry['laryng']

    drow = corr_table.add_row()
    style_cell(drow.cells[0], f"    {entry['airway']} (n\u2009f={entry['n_f']}, m={entry['n_m']})",
               size=8, align=WD_ALIGN_PARAGRAPH.LEFT)
    style_cell(drow.cells[1], entry['rho_f'], size=8)
    write_corr_p_cell(drow.cells[2], entry['p_f'], entry['q_f'])
    style_cell(drow.cells[3], entry['rho_m'], size=8)
    write_corr_p_cell(drow.cells[4], entry['p_m'], entry['q_m'])

for cell in corr_table.rows[-1].cells:
    set_cell_border(cell, bottom="single", bottom_sz="12", bottom_color="000000")

for row in corr_table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            pf = p.paragraph_format
            if pf.space_before is None:
                pf.space_before = Pt(1)
            if pf.space_after is None:
                pf.space_after = Pt(1)

corr_legend = doc.add_paragraph()
corr_legend.paragraph_format.space_before = Pt(6)
corr_run = corr_legend.add_run(
    'Spearman rank correlation coefficient (\u03c1) with corresponding P value. '
    'Computed separately for female and male participants within matched pairs. '
    'qFDR = Benjamini\u2013Hochberg adjusted P value across all comparisons in the '
    'same sex column. Asterisk (*) marks qFDR < 0.05; bold marks raw P < 0.05. '
    'Difficulty rated 0\u201310 (10 = most difficult); '
    'C/L = Cormack\u2013Lehane grade (1\u20134, 4 = poorest view).'
)
corr_run.font.size = Pt(7.5)
corr_run.font.name = 'Arial'
corr_run.italic = True

out_file = 'Tabelle_Schwierigkeit_Laryng_matched.docx'
doc.save(out_file)
print(f"Gespeichert: {out_file}")
