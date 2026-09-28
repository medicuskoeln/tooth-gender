import pandas as pd
import numpy as np
from scipy.stats import mannwhitneyu, shapiro, ttest_ind
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

# --- Data ---
df = pd.read_excel('roh1.xlsx')
proband_cols = ['PROBAND', 'Geschlecht_w1_m2', 'Ausbild_WA1_FA2_OA3',
                'Anz_SLT_4', 'Anz_DLT_4', 'Anz_MACI_4', 'Anz_GLIDE_CMAC_4', 'Anz_KingV_4',
                'NL_SLT_Easy_5', 'Nl_SLT_Diff_6', 'Nl_DLT_Easy_7', 'NI_DLT_Diff_8',
                'DA_SLT_Easy_9', 'DA_SLT_Diff_10', 'DA_DLT_Easy_11', 'DA_DLT_Diff_12',
                'Dent_Risk_Low_13', 'Dent_Risk_High_14', 'Dent_Damage_SLT_15', 'Dent_Damage_DLT_16']
df_prob = df[proband_cols].drop_duplicates(subset='PROBAND')
df_prob = df_prob[df_prob['Geschlecht_w1_m2'].isin([1, 2])]

# Groups
groups = [
    ('Total', df_prob),
    ('Residents', df_prob[df_prob['Ausbild_WA1_FA2_OA3'] == 1]),
    ('Specialists', df_prob[df_prob['Ausbild_WA1_FA2_OA3'] == 2]),
    ('Attendings', df_prob[df_prob['Ausbild_WA1_FA2_OA3'] == 3]),
]

for name, grp in groups:
    nf = (grp['Geschlecht_w1_m2'] == 1).sum()
    nm = (grp['Geschlecht_w1_m2'] == 2).sum()
    print(f"{name}: n={len(grp)} (female={nf}, male={nm})")

# --- Statistics ---
def compare_fm(col, grp):
    f = grp[grp['Geschlecht_w1_m2'] == 1]
    m = grp[grp['Geschlecht_w1_m2'] == 2]
    vf = pd.to_numeric(f[col], errors='coerce').dropna()
    vm = pd.to_numeric(m[col], errors='coerce').dropna()
    if len(vf) < 3 or len(vm) < 3:
        return '-', '-', '-'
    try:
        ps_f = shapiro(vf)[1]
        ps_m = shapiro(vm)[1]
    except:
        ps_f, ps_m = 0, 0
    if ps_f > 0.05 and ps_m > 0.05:
        _, p_val = ttest_ind(vf, vm)
        sf = f"{vf.mean():.1f} \u00b1 {vf.std():.1f}"
        sm = f"{vm.mean():.1f} \u00b1 {vm.std():.1f}"
    else:
        _, p_val = mannwhitneyu(vf, vm, alternative='two-sided')
        def fmt(v):
            iqr = v.quantile(0.75) - v.quantile(0.25)
            return f"{v.median():.1f} ({iqr:.1f})"
        sf = fmt(vf)
        sm = fmt(vm)
    p_str = f"{p_val:.3f}" if p_val >= 0.001 else "<0.001"
    return sf, sm, p_str

# --- Variables ---
cont_vars = [
    ('Anz_SLT_4', 'No. of SLT intubations'),
    ('Anz_DLT_4', 'No. of DLT intubations'),
    ('Anz_MACI_4', 'No. of Macintosh uses'),
    ('Anz_GLIDE_CMAC_4', 'No. of GlideScope/C-MAC uses'),
    ('Anz_KingV_4', 'No. of King Vision uses'),
    ('NL_SLT_Easy_5', 'NL SLT easy'),
    ('Nl_SLT_Diff_6', 'NL SLT difficult'),
    ('Nl_DLT_Easy_7', 'NL DLT easy'),
    ('NI_DLT_Diff_8', 'NL DLT difficult'),
    ('DA_SLT_Easy_9', 'DA SLT easy'),
    ('DA_SLT_Diff_10', 'DA SLT difficult'),
    ('DA_DLT_Easy_11', 'DA DLT easy'),
    ('DA_DLT_Diff_12', 'DA DLT difficult'),
    ('Dent_Risk_Low_13', 'Dental risk low'),
    ('Dent_Risk_High_14', 'Dental risk high'),
    ('Dent_Damage_SLT_15', 'Dental damage SLT'),
    ('Dent_Damage_DLT_16', 'Dental damage DLT'),
]

# Build data
table_data = []
for col, label in cont_vars:
    row_values = []
    for gname, grp in groups:
        sf, sm, p_str = compare_fm(col, grp)
        row_values.append({'f': sf, 'm': sm, 'p': p_str})
    table_data.append({'label': label, 'values': row_values})

# --- Word Document ---
doc = Document()

section = doc.sections[-1]
section.orientation = WD_ORIENT.LANDSCAPE
w, h = section.page_height, section.page_width
section.page_width = w
section.page_height = h
section.left_margin = Cm(1.2)
section.right_margin = Cm(1.2)
section.top_margin = Cm(1.5)

# Title
title = doc.add_heading('', level=1)
run = title.add_run(
    'Table 1. Demographic and Experience Characteristics of Participants '
    'Stratified by Training Level and Sex'
)
run.font.size = Pt(11)
run.font.name = 'Arial'
run.bold = True

# Legend
legend = doc.add_paragraph()
legend.paragraph_format.space_after = Pt(4)
run = legend.add_run(
    'Data are presented as mean \u00b1 SD or median (IQR) depending on distribution '
    '(Shapiro\u2013Wilk test). P values from independent-samples t-test or Mann\u2013Whitney U test. '
    'SLT = single-lumen tube; DLT = double-lumen tube; NL = normal laryngoscopy; '
    'DA = difficult airway; C-MAC = C-MAC videolaryngoscope.'
)
run.font.size = Pt(8)
run.font.name = 'Arial'
run.italic = True

# Table: 1 + 4*3 = 13 columns
ncols = 13
table = doc.add_table(rows=1, cols=ncols)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = 'Normal Table'

def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{kwargs.get("top","nil")}" w:sz="{kwargs.get("top_sz","0")}" w:space="0" w:color="{kwargs.get("top_color","000000")}"/>'
        f'<w:bottom w:val="{kwargs.get("bottom","nil")}" w:sz="{kwargs.get("bottom_sz","0")}" w:space="0" w:color="{kwargs.get("bottom_color","000000")}"/>'
        f'<w:left w:val="nil" w:sz="0" w:space="0" w:color="000000"/>'
        f'<w:right w:val="nil" w:sz="0" w:space="0" w:color="000000"/>'
        f'</w:tcBorders>')
    tcPr.append(tcBorders)

def style_cell(cell, text, bold=False, size=8, align=WD_ALIGN_PARAGRAPH.CENTER):
    cell.text = text
    for p in cell.paragraphs:
        p.alignment = align
        for r in p.runs:
            r.bold = bold
            r.font.size = Pt(size)
            r.font.name = 'Arial'

# Row 0: Group headers (merged)
r0 = table.rows[0]
style_cell(r0.cells[0], '', bold=True)
group_labels = ['Total', 'Residents', 'Specialists', 'Attendings']
for i, glabel in enumerate(group_labels):
    start = 1 + i * 3
    r0.cells[start].merge(r0.cells[start + 2])
    n_grp = len(groups[i][1])
    style_cell(r0.cells[start], f'{glabel}\n(n = {n_grp})', bold=True, size=8)

for cell in r0.cells:
    set_cell_border(cell, top="single", top_sz="12", top_color="000000")

# Row 1: Sub-headers Female / Male / P
r1 = table.add_row()
style_cell(r1.cells[0], 'Variable', bold=True, size=8, align=WD_ALIGN_PARAGRAPH.LEFT)
for i in range(4):
    start = 1 + i * 3
    nf = (groups[i][1]['Geschlecht_w1_m2'] == 1).sum()
    nm = (groups[i][1]['Geschlecht_w1_m2'] == 2).sum()
    style_cell(r1.cells[start], f'Female\n(n={nf})', bold=True, size=7)
    style_cell(r1.cells[start + 1], f'Male\n(n={nm})', bold=True, size=7)
    style_cell(r1.cells[start + 2], 'P', bold=True, size=7)

for cell in r1.cells:
    set_cell_border(cell, bottom="single", bottom_sz="8", bottom_color="000000")

# Data rows
for entry in table_data:
    row = table.add_row()
    style_cell(row.cells[0], entry['label'], size=8, align=WD_ALIGN_PARAGRAPH.LEFT)
    for i, vals in enumerate(entry['values']):
        start = 1 + i * 3
        style_cell(row.cells[start], vals['f'], size=7)
        style_cell(row.cells[start + 1], vals['m'], size=7)
        style_cell(row.cells[start + 2], vals['p'], size=7)
        try:
            pv = float(vals['p']) if vals['p'] != '<0.001' else 0.0001
            if pv < 0.05:
                for p_elem in row.cells[start + 2].paragraphs:
                    for r in p_elem.runs:
                        r.bold = True
        except:
            pass

# Bottom border
for cell in table.rows[-1].cells:
    set_cell_border(cell, bottom="single", bottom_sz="12", bottom_color="000000")

# Compact spacing
for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            pf = p.paragraph_format
            pf.space_before = Pt(1)
            pf.space_after = Pt(1)

doc.save('Tab1-oa.docx')
print("\nSaved: Tab1-oa.docx")
