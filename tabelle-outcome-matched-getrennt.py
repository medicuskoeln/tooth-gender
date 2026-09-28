import pandas as pd
import numpy as np
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

from stats_helpers import (
    bh_fdr, fmt_p, paired_test, mixed_effects_p, LEGEND_FDR_MEM,
)

# --- Daten einlesen ---
df = pd.read_excel('roh1.xlsx')
df = df[df['Geschlecht_w1_m2'].isin([1, 2])]

# WA/FA/OA bleiben getrennt
matching_cols = ['Ausbild_WA1_FA2_OA3', 'Laryng_diff_1Mac_2Glide_3King_4CMAC', 'tubus', 'atemweg']

matched_pairs = []
for _, group in df.groupby(matching_cols):
    weiblich = group[group['Geschlecht_w1_m2'] == 1]
    maennlich = group[group['Geschlecht_w1_m2'] == 2]
    n = min(len(weiblich), len(maennlich))
    if n > 0:
        w_sample = weiblich.sample(n=n, random_state=42).reset_index(drop=True)
        m_sample = maennlich.sample(n=n, random_state=42).reset_index(drop=True)
        matched_pairs.append((w_sample, m_sample))

df_w_all = pd.concat([w for w, _ in matched_pairs], ignore_index=True)
df_m_all = pd.concat([m for _, m in matched_pairs], ignore_index=True)

print(f"Gematchte Paare gesamt: {len(df_w_all)}")
print(f"  Standard Airway: w={len(df_w_all[df_w_all['atemweg']==1])}, m={len(df_m_all[df_m_all['atemweg']==1])}")
print(f"  Difficult Airway: w={len(df_w_all[df_w_all['atemweg']==2])}, m={len(df_m_all[df_m_all['atemweg']==2])}")

outcome_vars = [
    ('res_max_alle_zaehne', 'Force all front teeth'),
    ('res_12_max', 'Tooth 12'),
    ('res_11_max', 'Tooth 11'),
    ('res_21_max', 'Tooth 21'),
    ('res_22_max', 'Tooth 22'),
    ('BURP', 'BURP'),
    ('Intubationsdauer', 'Intubation duration'),
    ('Sicht_Cormack', 'C&L grade'),
]


def analyse_matched(col, df_w, df_m):
    vals_w = pd.to_numeric(df_w[col], errors='coerce')
    vals_m = pd.to_numeric(df_m[col], errors='coerce')
    mask = (~vals_w.isna()) & (~vals_m.isna())
    vw = vals_w[mask]
    vm = vals_m[mask]
    if len(vw) < 3:
        return '-', '-', np.nan, np.nan, 0
    diff = vw.values - vm.values
    from scipy.stats import shapiro
    try:
        p_shap = shapiro(diff)[1]
    except Exception:
        p_shap = 0
    if p_shap > 0.05:
        sw = f"{vw.mean():.2f} \u00b1 {vw.std():.2f}"
        sm = f"{vm.mean():.2f} \u00b1 {vm.std():.2f}"
    else:
        def fmt_median(v):
            return (f"{v.median():.2f} ({v.quantile(0.25):.2f}\u2013{v.quantile(0.75):.2f}) "
                    f"[{v.min():.2f}\u2013{v.max():.2f}]")
        sw = fmt_median(vw)
        sm = fmt_median(vm)
    p_raw, _ = paired_test(vw.values, vm.values)
    p_mm = mixed_effects_p(df_w.loc[mask.values], df_m.loc[mask.values], col)
    return sw, sm, p_raw, p_mm, len(vw)


results = []
all_p_raw = []
all_p_mm = []
for col, label in outcome_vars:
    row_data = {'label': label}
    for aw, aw_label in [(1, 'standard'), (2, 'difficult')]:
        w_sub = df_w_all[df_w_all['atemweg'] == aw].reset_index(drop=True)
        m_sub = df_m_all[df_m_all['atemweg'] == aw].reset_index(drop=True)
        sw, sm, p_raw, p_mm, n = analyse_matched(col, w_sub, m_sub)
        row_data[f'{aw_label}_w'] = sw
        row_data[f'{aw_label}_m'] = sm
        row_data[f'{aw_label}_p'] = p_raw
        row_data[f'{aw_label}_pmm'] = p_mm
        all_p_raw.append(p_raw)
        all_p_mm.append(p_mm)
    results.append(row_data)

q_all = bh_fdr(all_p_raw)
qmm_all = bh_fdr(all_p_mm)
idx = 0
for r in results:
    for aw_label in ('standard', 'difficult'):
        r[f'{aw_label}_q'] = q_all[idx]
        r[f'{aw_label}_qmm'] = qmm_all[idx]
        idx += 1

n_std_w = len(df_w_all[df_w_all['atemweg'] == 1])
n_std_m = len(df_m_all[df_m_all['atemweg'] == 1])
n_dif_w = len(df_w_all[df_w_all['atemweg'] == 2])
n_dif_m = len(df_m_all[df_m_all['atemweg'] == 2])

# ============================================================
# Word-Dokument
# ============================================================
doc = Document()
section = doc.sections[-1]
section.orientation = WD_ORIENT.LANDSCAPE
new_width, new_height = section.page_height, section.page_width
section.page_width = new_width
section.page_height = new_height
section.left_margin = Cm(1.5)
section.right_margin = Cm(1.5)

doc.add_heading('Table: Matched-Pair Outcome Variables by Airway Difficulty and Sex (WA/FA/OA getrennt)', level=1)
doc.add_paragraph(
    "Matched-pair analysis. WA, FA, and OA kept as separate education levels. "
    "Pairs matched on: education level (WA/FA/OA), laryngoscope type, tube type, airway difficulty.\n"
    "Data presented as median (IQR) [range] or mean \u00b1 SD depending on distribution "
    "(Shapiro-Wilk on paired differences).\n"
    f"Total matched pairs: {len(df_w_all)} (Standard: {n_std_w}, Difficult: {n_dif_w})"
)

table = doc.add_table(rows=1, cols=7)
table.alignment = WD_TABLE_ALIGNMENT.CENTER

header_row = table.rows[0]
header_row.cells[0].text = ''
header_row.cells[1].merge(header_row.cells[3])
header_row.cells[1].text = 'Standard Airway'
header_row.cells[4].merge(header_row.cells[6])
header_row.cells[4].text = 'Difficult Airway'

for cell in header_row.cells:
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(9)
            run.font.name = 'Arial'

sub_row = table.add_row()
sub_headers = ['Variable',
               f'Female\n(n={n_std_w})', f'Male\n(n={n_std_m})', 'P / qFDR / pMM',
               f'Female\n(n={n_dif_w})', f'Male\n(n={n_dif_m})', 'P / qFDR / pMM']
for i, txt in enumerate(sub_headers):
    sub_row.cells[i].text = txt
    for p in sub_row.cells[i].paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(8)
            run.font.name = 'Arial'


def write_p_cell(cell, p_raw, p_fdr, p_mm):
    cell.text = ''
    para = cell.paragraphs[0]
    para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para.paragraph_format.space_before = Pt(0)
    para.paragraph_format.space_after = Pt(0)
    sig = ''
    try:
        if p_fdr is not None and not np.isnan(p_fdr) and p_fdr < 0.05:
            sig += '*'
    except Exception:
        pass
    try:
        if p_mm is not None and not np.isnan(p_mm) and p_mm < 0.05:
            sig += '\u2020'
    except Exception:
        pass
    r1 = para.add_run(fmt_p(p_raw) + sig)
    r1.font.size = Pt(8)
    r1.font.name = 'Arial'
    try:
        if p_raw is not None and not np.isnan(float(p_raw)) and float(p_raw) < 0.05:
            r1.bold = True
    except Exception:
        pass
    para2 = cell.add_paragraph()
    para2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para2.paragraph_format.space_before = Pt(0)
    para2.paragraph_format.space_after = Pt(0)
    r2 = para2.add_run(f"qFDR={fmt_p(p_fdr)}")
    r2.font.size = Pt(6.5)
    r2.font.name = 'Arial'
    r2.italic = True
    para3 = cell.add_paragraph()
    para3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    para3.paragraph_format.space_before = Pt(0)
    para3.paragraph_format.space_after = Pt(0)
    r3 = para3.add_run(f"pMM={fmt_p(p_mm)}")
    r3.font.size = Pt(6.5)
    r3.font.name = 'Arial'
    r3.italic = True


for r in results:
    data_row = table.add_row()
    data_row.cells[0].text = r['label']
    for p in data_row.cells[0].paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
        for run in p.runs:
            run.font.size = Pt(8)
            run.font.name = 'Arial'
    data_row.cells[1].text = r['standard_w']
    data_row.cells[2].text = r['standard_m']
    write_p_cell(data_row.cells[3], r['standard_p'], r['standard_q'], r['standard_pmm'])
    data_row.cells[4].text = r['difficult_w']
    data_row.cells[5].text = r['difficult_m']
    write_p_cell(data_row.cells[6], r['difficult_p'], r['difficult_q'], r['difficult_pmm'])
    for i in (1, 2, 4, 5):
        for p in data_row.cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8)
                run.font.name = 'Arial'


def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{kwargs.get("top", "nil")}" w:sz="{kwargs.get("top_sz", "0")}" w:space="0" w:color="{kwargs.get("top_color", "000000")}"/>'
        f'<w:bottom w:val="{kwargs.get("bottom", "nil")}" w:sz="{kwargs.get("bottom_sz", "0")}" w:space="0" w:color="{kwargs.get("bottom_color", "000000")}"/>'
        f'<w:left w:val="nil" w:sz="0" w:space="0" w:color="000000"/>'
        f'<w:right w:val="nil" w:sz="0" w:space="0" w:color="000000"/>'
        f'</w:tcBorders>')
    tcPr.append(tcBorders)


table.style = 'Normal Table'
for cell in table.rows[0].cells:
    set_cell_border(cell, top="single", top_sz="12", top_color="000000")
for cell in table.rows[1].cells:
    set_cell_border(cell, bottom="single", bottom_sz="8", bottom_color="000000")
for cell in table.rows[-1].cells:
    set_cell_border(cell, bottom="single", bottom_sz="12", bottom_color="000000")

legend_p = doc.add_paragraph()
legend_p.paragraph_format.space_before = Pt(6)
legend_run = legend_p.add_run(LEGEND_FDR_MEM)
legend_run.font.size = Pt(7.5)
legend_run.font.name = 'Arial'
legend_run.italic = True

doc.save('Tabelle_Outcome_Matched_getrennt.docx')
print("\nTabelle gespeichert in Tabelle_Outcome_Matched_getrennt.docx")

for r in results:
    print(f"{r['label']}: std p={fmt_p(r['standard_p'])} q={fmt_p(r['standard_q'])} pMM={fmt_p(r['standard_pmm'])} | "
          f"diff p={fmt_p(r['difficult_p'])} q={fmt_p(r['difficult_q'])} pMM={fmt_p(r['difficult_pmm'])}")
