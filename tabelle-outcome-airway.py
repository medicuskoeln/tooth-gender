import pandas as pd
import numpy as np
from scipy.stats import mannwhitneyu, shapiro, ttest_ind
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

# --- Daten einlesen ---
df = pd.read_excel('roh1.xlsx')

# Geschlecht bereinigen (Leerzeichen entfernen)
df = df[df['Geschlecht_w1_m2'].isin([1, 2])]

# Variablen definieren
outcome_vars = [
    ('res_max_alle_zaehne', 'Force all front teeth'),
    ('res_12_max', 'Tooth 12'),
    ('res_11_max', 'Tooth 11'),
    ('res_21_max', 'Tooth 21'),
    ('res_22_max', 'Tooth 22'),
    ('BURP', 'BURP'),
    ('Intubationsdauer', 'Intubation duration'),
    ('Sicht_Cormack', 'C&L grade'),
]

# --- Statistische Analyse ---
def analyse_group(vals_w, vals_m):
    """Vergleich w vs m: gibt Statistik-Strings und p-Wert zurück."""
    vals_w = pd.to_numeric(vals_w, errors='coerce').dropna()
    vals_m = pd.to_numeric(vals_m, errors='coerce').dropna()
    if len(vals_w) < 3 or len(vals_m) < 3:
        return '-', '-', '-', '-', 0

    # Normalverteilung prüfen
    try:
        p_w = shapiro(vals_w)[1]
        p_m = shapiro(vals_m)[1]
    except:
        p_w, p_m = 0, 0

    if p_w > 0.05 and p_m > 0.05:
        _, p_val = ttest_ind(vals_w, vals_m)
        stat_w = f"{vals_w.mean():.2f} ± {vals_w.std():.2f}"
        stat_m = f"{vals_m.mean():.2f} ± {vals_m.std():.2f}"
        fmt = "mean_sd"
    else:
        _, p_val = mannwhitneyu(vals_w, vals_m, alternative='two-sided')
        # Median (IQR) [Range]
        def fmt_median(v):
            med = v.median()
            q25 = v.quantile(0.25)
            q75 = v.quantile(0.75)
            rng_min = v.min()
            rng_max = v.max()
            return f"{med:.2f} ({q25:.2f}–{q75:.2f}) [{rng_min:.2f}–{rng_max:.2f}]"
        stat_w = fmt_median(vals_w)
        stat_m = fmt_median(vals_m)
        fmt = "median_iqr"

    p_str = f"{p_val:.3f}" if p_val >= 0.001 else "<0.001"
    return stat_w, stat_m, p_str, fmt, len(vals_w) + len(vals_m)

# Auswertung für Standard (atemweg=1) und Difficult (atemweg=2)
results = []
for col, label in outcome_vars:
    row_data = {'label': label}
    for aw, aw_label in [(1, 'standard'), (2, 'difficult')]:
        sub = df[df['atemweg'] == aw]
        w = sub[sub['Geschlecht_w1_m2'] == 1][col]
        m = sub[sub['Geschlecht_w1_m2'] == 2][col]
        stat_w, stat_m, p_str, fmt, n = analyse_group(w, m)
        row_data[f'{aw_label}_w'] = stat_w
        row_data[f'{aw_label}_m'] = stat_m
        row_data[f'{aw_label}_p'] = p_str
        row_data[f'{aw_label}_fmt'] = fmt
    results.append(row_data)

# Stichprobengrößen
n_std_w = len(df[(df['atemweg'] == 1) & (df['Geschlecht_w1_m2'] == 1)])
n_std_m = len(df[(df['atemweg'] == 1) & (df['Geschlecht_w1_m2'] == 2)])
n_dif_w = len(df[(df['atemweg'] == 2) & (df['Geschlecht_w1_m2'] == 1)])
n_dif_m = len(df[(df['atemweg'] == 2) & (df['Geschlecht_w1_m2'] == 2)])

# --- Word-Dokument im Anaesthesiology-Stil ---
doc = Document()

# Querformat
section = doc.sections[-1]
section.orientation = WD_ORIENT.LANDSCAPE
new_width, new_height = section.page_height, section.page_width
section.page_width = new_width
section.page_height = new_height
section.left_margin = Cm(1.5)
section.right_margin = Cm(1.5)

doc.add_heading('Table: Outcome Variables by Airway Difficulty and Sex', level=1)

# Fußnoten-Text
doc.add_paragraph(
    "Data presented as median (IQR) [range] or mean ± SD depending on distribution (Shapiro-Wilk test). "
    "P-values from Mann-Whitney U test or independent t-test."
)

# Tabelle erstellen: 7 Spalten
# | Variable | Female | Male | P | Female | Male | P |
table = doc.add_table(rows=1, cols=7)
table.alignment = WD_TABLE_ALIGNMENT.CENTER

# --- Obere Überschriftszeile (merged): Standard Airway / Difficult Airway ---
header_row = table.rows[0]
# Variable-Zelle
header_row.cells[0].text = ''
# Standard Airway (Spalten 1-3 mergen)
header_row.cells[1].merge(header_row.cells[3])
header_row.cells[1].text = 'Standard Airway'
# Difficult Airway (Spalten 4-6 mergen)
header_row.cells[4].merge(header_row.cells[6])
header_row.cells[4].text = 'Difficult Airway'

# Zentrieren und fett
for cell in header_row.cells:
    for p in cell.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(9)
            run.font.name = 'Arial'

# --- Zweite Zeile: Sub-Header ---
sub_row = table.add_row()
sub_headers = ['Variable',
               f'Female\n(n={n_std_w})', f'Male\n(n={n_std_m})', 'P value',
               f'Female\n(n={n_dif_w})', f'Male\n(n={n_dif_m})', 'P value']
for i, txt in enumerate(sub_headers):
    sub_row.cells[i].text = txt
    for p in sub_row.cells[i].paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.bold = True
            run.font.size = Pt(8)
            run.font.name = 'Arial'

# --- Datenzeilen ---
for r in results:
    data_row = table.add_row()
    vals = [r['label'],
            r['standard_w'], r['standard_m'], r['standard_p'],
            r['difficult_w'], r['difficult_m'], r['difficult_p']]
    for i, val in enumerate(vals):
        data_row.cells[i].text = str(val)
        for p in data_row.cells[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if i > 0 else WD_ALIGN_PARAGRAPH.LEFT
            for run in p.runs:
                run.font.size = Pt(8)
                run.font.name = 'Arial'
                # p-Wert fett wenn signifikant
                if i in [3, 6] and val not in ['-', '']:
                    try:
                        pv = float(val) if val != '<0.001' else 0.0001
                        if pv < 0.05:
                            run.bold = True
                    except:
                        pass

# --- Anaesthesiology-Stil: Nur horizontale Linien ---
# Alle Ränder entfernen, dann nur top/bottom-Linien setzen
def set_cell_border(cell, **kwargs):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
        f'<w:top w:val="{kwargs.get("top", "nil")}" w:sz="{kwargs.get("top_sz", "0")}" w:space="0" w:color="{kwargs.get("top_color", "000000")}"/>'
        f'<w:bottom w:val="{kwargs.get("bottom", "nil")}" w:sz="{kwargs.get("bottom_sz", "0")}" w:space="0" w:color="{kwargs.get("bottom_color", "000000")}"/>'
        f'<w:left w:val="nil" w:sz="0" w:space="0" w:color="000000"/>'
        f'<w:right w:val="nil" w:sz="0" w:space="0" w:color="000000"/>'
        f'</w:tcBorders>')
    tcPr.append(tcBorders)

# Entferne den Standard Table Style
table.style = 'Normal Table'

# Obere Linie (über Header)
for cell in table.rows[0].cells:
    set_cell_border(cell, top="single", top_sz="12", top_color="000000")

# Linie unter Sub-Header
for cell in table.rows[1].cells:
    set_cell_border(cell, bottom="single", bottom_sz="8", bottom_color="000000")

# Linie unter letzter Datenzeile
last_row = table.rows[-1]
for cell in last_row.cells:
    set_cell_border(cell, bottom="single", bottom_sz="12", bottom_color="000000")

# Zellenabstand minimieren
for row in table.rows:
    for cell in row.cells:
        for p in cell.paragraphs:
            pf = p.paragraph_format
            pf.space_before = Pt(1)
            pf.space_after = Pt(1)

doc.save('Tabelle_Outcome_Airway_Sex.docx')
print("Tabelle gespeichert in Tabelle_Outcome_Airway_Sex.docx")

# Ergebnisse auch in Konsole ausgeben
print(f"\nStandard Airway: female n={n_std_w}, male n={n_std_m}")
print(f"Difficult Airway: female n={n_dif_w}, male n={n_dif_m}\n")
for r in results:
    print(f"{r['label']}:")
    print(f"  Standard: w={r['standard_w']}  m={r['standard_m']}  p={r['standard_p']}")
    print(f"  Difficult: w={r['difficult_w']}  m={r['difficult_m']}  p={r['difficult_p']}")
