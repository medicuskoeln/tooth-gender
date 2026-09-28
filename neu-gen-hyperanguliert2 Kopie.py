import pandas as pd
import numpy as np
from scipy.stats import ttest_rel, wilcoxon, shapiro
from docx import Document
from docx.shared import Inches
from docx.enum.section import WD_ORIENT
import matplotlib.pyplot as plt
import os

def analyse_matched_pairs(df, analyse_cols, matching_cols, hyper_wert, doc_filename):
    # Filter nach hyperanguliert, nur wenn hyper_wert angegeben ist
    if hyper_wert is not None:
        df = df[df['hyperanguliert'] == hyper_wert]

    # Paare bilden
    matched_pairs = []
    for _, group in df.groupby(matching_cols):
        weiblich = group[group['Geschlecht_w1_m2'] == 1]
        maennlich = group[group['Geschlecht_w1_m2'] == 2]
        n = min(len(weiblich), len(maennlich))
        if n > 0:
            w_sample = weiblich.sample(n=n, random_state=42).reset_index(drop=True)
            m_sample = maennlich.sample(n=n, random_state=42).reset_index(drop=True)
            matched_pairs.append((w_sample, m_sample))

    # Word-Dokument vorbereiten
    doc = Document()

    # Einleitungstext für das Word-Dokument
    einleitung = (
        "In dieser Matched-Pair-Analyse wurden Paare gebildet, die in folgenden Parametern übereinstimmen:\n"
        "- Ausbild_WA1_FA2_OA3\n"
        "- hyperanguliert" + (f" = {hyper_wert} (0=nein, 1=ja)" if hyper_wert is not None else " (alle Werte)")
        + "\n- tubus\n- atemweg\n\n"
        "Für jedes Paar wurde jeweils ein Datensatz mit Geschlecht = weiblich und einer mit Geschlecht = männlich ausgewählt.\n\n"
        "Untersucht wurden folgende abhängige Variablen:\n"
        + "\n".join(f"- {v}" for v in analyse_cols)
    )
    doc.add_heading('Matched-Pair-Analyse mit statistischer Auswertung', 0)
    doc.add_paragraph(einleitung)
    doc.add_paragraph(f"Gesamtanzahl an Datensätzen: {len(df)}")
    doc.add_paragraph(f"Anzahl der gebildeten Paare: {sum(len(w) for w, _ in matched_pairs)}")

    # Übersichtstabelle vorbereiten
    results_table = []

    # Plot-Verzeichnis
    os.makedirs('plots', exist_ok=True)

    # Analyse je Variable
    for res in analyse_cols:
        weiblich_vals, maennlich_vals = [], []

        for w, m in matched_pairs:
            if res in w.columns and res in m.columns:
                w_vals = w[res]
                m_vals = m[res]
                if len(w_vals) == len(m_vals):
                    weiblich_vals.extend(w_vals.tolist())
                    maennlich_vals.extend(m_vals.tolist())

        if len(weiblich_vals) < 3:
            continue

        # In numerisch umwandeln & NaN entfernen
        weiblich_vals = pd.to_numeric(pd.Series(weiblich_vals), errors='coerce')
        maennlich_vals = pd.to_numeric(pd.Series(maennlich_vals), errors='coerce')
        mask = (~weiblich_vals.isna()) & (~maennlich_vals.isna())
        weiblich_vals = weiblich_vals[mask]
        maennlich_vals = maennlich_vals[mask]

        if len(weiblich_vals) < 3:
            continue

        # Normalverteilung prüfen
        differences = weiblich_vals - maennlich_vals
        p_shapiro = shapiro(differences)[1]

        # Statistische Kennwerte berechnen
        if p_shapiro > 0.05:
            test_name = "t-Test"
            t_stat, p_val = ttest_rel(weiblich_vals, maennlich_vals, nan_policy='omit')
            mw_w = weiblich_vals.mean()
            std_w = weiblich_vals.std()
            mw_m = maennlich_vals.mean()
            std_m = maennlich_vals.std()
            kennwerte_text = (
                f"Weiblich: Mittelwert = {mw_w:.2f}, SD = {std_w:.2f}\n"
                f"Männlich: Mittelwert = {mw_m:.2f}, SD = {std_m:.2f}"
            )
        else:
            test_name = "Wilcoxon-Test"
            try:
                t_stat, p_val = wilcoxon(weiblich_vals, maennlich_vals)
            except:
                continue
            med_w = weiblich_vals.median()
            iqr_w = weiblich_vals.quantile(0.75) - weiblich_vals.quantile(0.25)
            range_w = weiblich_vals.max() - weiblich_vals.min()
            med_m = maennlich_vals.median()
            iqr_m = maennlich_vals.quantile(0.75) - maennlich_vals.quantile(0.25)
            range_m = maennlich_vals.max() - maennlich_vals.min()
            kennwerte_text = (
                f"Weiblich: Median = {med_w:.2f}, IQR = {iqr_w:.2f}, Range = {range_w:.2f}\n"
                f"Männlich: Median = {med_m:.2f}, IQR = {iqr_m:.2f}, Range = {range_m:.2f}"
            )

        effect_size = cohens_d(weiblich_vals, maennlich_vals)

        # Ergebnis in Tabelle speichern
        results_table.append({
            'Variable': res,
            'Test': test_name,
            'p-Wert': round(p_val, 4),
            'Signifikant': "Ja" if p_val < 0.05 else "Nein",
            'Effektstärke': round(effect_size, 3),
            'weiblich_vals': weiblich_vals.tolist(),
            'maennlich_vals': maennlich_vals.tolist()
        })

        # Boxplot
        plt.figure(figsize=(5, 4))
        plt.boxplot([weiblich_vals, maennlich_vals], tick_labels=['weiblich', 'männlich'])
        plt.title(f'{res} – {test_name}\np={p_val:.4f}')
        plt.ylabel('Wert')
        plt.tight_layout()
        box_path = f'plots/box_{res}_laryng{hyper_wert}.png'
        plt.savefig(box_path)
        plt.close()

        # Histogramm
        plt.figure(figsize=(5, 4))
        plt.hist(weiblich_vals, bins=10, alpha=0.6, label='weiblich')
        plt.hist(maennlich_vals, bins=10, alpha=0.6, label='männlich')
        plt.title(f'{res} – Histogramm')
        plt.xlabel('Wert')
        plt.ylabel('Häufigkeit')
        plt.legend()
        plt.tight_layout()
        hist_path = f'plots/hist_{res}_laryng{hyper_wert}.png'
        plt.savefig(hist_path)
        plt.close()

        # Ergebnisse + Plots ins Word-Dokument
        doc.add_heading(res, level=2)
        doc.add_paragraph(kennwerte_text)
        doc.add_paragraph(f"{test_name}: p = {p_val:.4f}, Effektstärke d = {effect_size:.3f}")
        doc.add_picture(box_path, width=Inches(4.5))
        doc.add_paragraph()
        doc.add_picture(hist_path, width=Inches(4.5))
        doc.add_paragraph()

    # Querformat für die Seite einstellen
    section = doc.sections[-1]
    section.orientation = WD_ORIENT.LANDSCAPE
    new_width, new_height = section.page_height, section.page_width
    section.page_width = new_width
    section.page_height = new_height

    # Übersichtstabelle einfügen
    doc.add_heading('Zusammenfassung der Ergebnisse', level=1)
    table = doc.add_table(rows=1, cols=9)
    hdr = table.rows[0].cells
    hdr[0].text = 'Variable'
    hdr[1].text = 'Test'
    hdr[2].text = 'p-Wert'
    hdr[3].text = 'Signifikant'
    hdr[4].text = 'Effektstärke (d)'
    hdr[5].text = 'Median (w/m)'
    hdr[6].text = 'Mittelwert (w/m)'
    hdr[7].text = 'SD (w/m)'
    hdr[8].text = 'IQR (w/m)'

    for res in results_table:
        row = table.add_row().cells
        row[0].text = res['Variable']
        row[1].text = res['Test']
        row[2].text = str(res['p-Wert'])
        row[3].text = res['Signifikant']
        row[4].text = str(res['Effektstärke'])
        w_vals = pd.to_numeric(pd.Series(res.get('weiblich_vals', [])), errors='coerce')
        m_vals = pd.to_numeric(pd.Series(res.get('maennlich_vals', [])), errors='coerce')
        if res['Test'] == "t-Test":
            # Nur Mittelwert und SD ausgeben
            row[5].text = "-"
            row[6].text = f"{w_vals.mean():.2f} / {m_vals.mean():.2f}"
            row[7].text = f"{w_vals.std():.2f} / {m_vals.std():.2f}"
            row[8].text = "-"
        else:
            # Nur Median und IQR ausgeben
            row[5].text = f"{w_vals.median():.2f} / {m_vals.median():.2f}"
            row[6].text = "-"
            row[7].text = "-"
            iqr_w = w_vals.quantile(0.75) - w_vals.quantile(0.25)
            iqr_m = m_vals.quantile(0.75) - m_vals.quantile(0.25)
            row[8].text = f"{iqr_w:.2f} / {iqr_m:.2f}"

    # Word-Dokument speichern
    doc.save(doc_filename)
    print(f"Analyse für hyperanguliert={hyper_wert} abgeschlossen und in {doc_filename} gespeichert.")
    if doc_filename == "matched_pair_analyse_gesamt.docx" and os.path.exists('plots/res_w_m_vergleich.png'):
        doc.add_picture('plots/res_w_m_vergleich.png', width=Inches(5))
    return results_table

# --- Hauptprogramm ---
# Daten einlesen
df = pd.read_excel('roh1.xlsx')
df = df[df['tubus'] == 2]  # <--- Nur tubus = 2 auswerten
comb_cols = [col for col in df.columns if col.startswith('comb')]
res_cols = [col for col in df.columns if col.startswith('res')]
extra_cols = ['BURP', 'Intubationsdauer', 'Sicht_Cormack']
analyse_cols = comb_cols + res_cols + [col for col in extra_cols if col in df.columns]
matching_cols = ['Ausbild_WA1_FA2_OA3', 'hyperanguliert', 'tubus', 'atemweg']

# Effektgröße berechnen (wie gehabt)
def cohens_d(x, y):
    x, y = np.array(x), np.array(y)
    diff = x - y
    if np.std(diff, ddof=1) == 0:
        return 0
    return np.mean(diff) / np.std(diff, ddof=1)

# Analyse für hyperanguliert=0 (nein)
analyse_matched_pairs(df, analyse_cols, matching_cols, hyper_wert=0, doc_filename="matched_pair_analyse_hyperanguliert0_nein.docx")

# Analyse für hyperanguliert=1 (ja)
analyse_matched_pairs(df, analyse_cols, matching_cols, hyper_wert=1, doc_filename="matched_pair_analyse_hyperanguliert1_ja.docx")

# Gesamtauswertung (ohne Filter auf hyperanguliert)
results_table_gesamt = analyse_matched_pairs(
    df, analyse_cols, matching_cols, hyper_wert=None, doc_filename="matched_pair_analyse_gesamt.docx"
)

# Sammle p-Werte für alle res-Variablen aus der Gesamtauswertung
res_pvals = []
res_names = []
for res in results_table_gesamt:
    if res['Variable'].startswith('res'):
        res_names.append(res['Variable'])
        res_pvals.append(res['p-Wert'])

if res_pvals:
    plt.figure(figsize=(max(8, len(res_names)*0.5), 5))
    bars = plt.bar(res_names, res_pvals, color=['red' if p < 0.05 else 'grey' for p in res_pvals])
    plt.axhline(0.05, color='blue', linestyle='--', label='p = 0.05')
    plt.ylabel('p-Wert')
    plt.title('Signifikanzwerte für res-Variablen')
    plt.xticks(rotation=90)
    plt.legend()
    plt.tight_layout()
    plt.savefig('plots/signifikanz_res_variablen.png')
    plt.close()

# Bild ggf. in ein Dokument einfügen, falls gewünscht
# doc.add_picture('plots/signifikanz_res_variablen.png', width=Inches(5))

# Diagramm: res-Werte für w/m nebeneinander, p-Wert falls signifikant
labels = []
w_vals = []
m_vals = []
p_vals = []
tests = []

for res in results_table_gesamt:
    if res['Variable'].startswith('res'):
        labels.append(res['Variable'])
        weiblich = pd.to_numeric(pd.Series(res.get('weiblich_vals', [])), errors='coerce')
        maennlich = pd.to_numeric(pd.Series(res.get('maennlich_vals', [])), errors='coerce')
        if res['Test'] == "t-Test":
            w_vals.append(weiblich.mean())
            m_vals.append(maennlich.mean())
        else:
            w_vals.append(weiblich.median())
            m_vals.append(maennlich.median())
        p_vals.append(res['p-Wert'])
        tests.append(res['Test'])

x = np.arange(len(labels))
width = 0.35

plt.figure(figsize=(max(8, len(labels)*0.7), 6))
bar1 = plt.bar(x - width/2, w_vals, width, label='weiblich')
bar2 = plt.bar(x + width/2, m_vals, width, label='männlich')

# p-Werte über signifikanten Balken anzeigen
for i, p in enumerate(p_vals):
    if p < 0.05:
        y = max(w_vals[i], m_vals[i])
        plt.text(x[i], y + 0.05*abs(y), f"p={p:.3f}", ha='center', va='bottom', color='red', fontsize=10, fontweight='bold')

plt.xticks(x, labels, rotation=90)
plt.ylabel('Mittelwert/Median')
plt.title('res-Variablen: Vergleich weiblich vs. männlich\n(Mittelwert bei t-Test, Median bei Wilcoxon)')
plt.legend()
plt.tight_layout()
plt.savefig('plots/res_w_m_vergleich.png')
plt.close()

# Nach der Ergebnistabelle im Gesamtdokument:
# doc.add_picture('plots/res_w_m_vergleich.png', width=Inches(5))

# Diagramm für Sicht_Cormack (wenn vorhanden)
for var in ['Sicht_Cormack', 'Intubationsdauer']:
    w_val = None
    m_val = None
    p_val = None
    test = None
    for res in results_table_gesamt:
        if res['Variable'] == var:
            weiblich = pd.to_numeric(pd.Series(res.get('weiblich_vals', [])), errors='coerce')
            maennlich = pd.to_numeric(pd.Series(res.get('maennlich_vals', [])), errors='coerce')
            if res['Test'] == "t-Test":
                w_val = weiblich.mean()
                m_val = maennlich.mean()
            else:
                w_val = weiblich.median()
                m_val = maennlich.median()
            p_val = res['p-Wert']
            test = res['Test']
            break
    if w_val is not None and m_val is not None:
        plt.figure(figsize=(4, 6))
        bars = plt.bar(['weiblich', 'männlich'], [w_val, m_val], color=['tab:blue', 'tab:orange'])
        if p_val is not None and p_val < 0.05:
            y = max(w_val, m_val)
            plt.text(0.5, y + 0.05*abs(y), f"p={p_val:.3f}", ha='center', va='bottom', color='red', fontsize=12, fontweight='bold')
        plt.ylabel('Mittelwert/Median')
        plt.title(f'{var}: Vergleich weiblich vs. männlich\n(Mittelwert bei t-Test, Median bei Wilcoxon)')
        plt.tight_layout()
        plt.savefig(f'plots/{var}_w_m_vergleich.png')
        plt.close()
